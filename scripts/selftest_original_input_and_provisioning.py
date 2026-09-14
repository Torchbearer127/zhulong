#!/usr/bin/env python3
"""Offline regressions; synthetic assets are not vulnerability or recording proof."""
import json
from pathlib import Path
import tempfile
import unittest
import importlib
import os
import subprocess
import sys

import validate_report_bundle as validator


class OriginalInputTests(unittest.TestCase):
    def test_renderer_embeds_declared_original_input_and_helper_displays_text(self):
        import render_confirmed_vuln_docx as renderer
        import selftest_plugin as fixtures
        from docx import Document
        from PIL import Image
        helper = importlib.import_module("original_input_evidence")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "attachments").mkdir()
            (root / "attachments/input.txt").write_text("synthetic original input\n")
            Image.new("RGB", (80, 40), "white").save(root / "attachments/input.png")
            finding = fixtures.issue23_finding(["docker compose up -d app"], include_history=True)
            finding["original_input"] = {"text_path": "attachments/input.txt", "screenshot_path": "attachments/input.png", "title": "PoC 原始输入", "description": "仅供测试。", "caption": "图 1：原始输入截图"}
            shell = renderer.build_generated_recording_shell(finding, "zh-CN", {}, {})
            display = "\n".join(validator.extract_shell_function_body(shell, "show_original_input"))
            result = subprocess.run(["sh", "-c", 'set -eu\nREPLAY_LOG=runtime.log\nPAUSE_LONG=0\npause_step() { :; }\n' + display], cwd=root, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("synthetic original input\n", result.stdout)
            self.assertEqual((root / "runtime.log").read_text(), "synthetic original input\n")
            manifest = helper.prepare_original_input(root, finding["original_input"], "zh-CN")
            (root / "findings.json").write_text(json.dumps({"findings": [{"original_input": finding["original_input"]}]}))
            doc = Document()
            helper.render_original_input(doc, root, manifest)
            doc.save(root / "report.docx")
            paragraphs = doc.paragraphs
            paragraphs[-1].text = "图 9：无关图片"
            doc.save(root / "report.docx")
            with self.assertRaisesRegex(ValueError, "DOCX"):
                helper.validate_original_input(root, required=True)

    def test_declared_input_requires_existing_screenshot_and_preserves_bytes(self):
        from docx import Document
        from PIL import Image
        helper = importlib.import_module("original_input_evidence")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "attachments").mkdir()
            raw = root / "attachments/input.txt"
            raw.write_bytes(b"GET /synthetic-fixture HTTP/1.1\r\nHost: example.invalid\r\n\r\n")
            spec = {"text_path": "attachments/input.txt", "screenshot_path": "attachments/input.png", "title": "PoC 原始输入", "description": "本图为独立测试图片，不是真实漏洞证据。", "caption": "图 1：原始输入截图"}
            with self.assertRaisesRegex(ValueError, "ORIGINAL_INPUT_INVALID"):
                helper.prepare_original_input(root, spec, "zh-CN")
            Image.new("RGB", (80, 40), "white").save(root / spec["screenshot_path"])
            before = (root / spec["screenshot_path"]).read_bytes()
            manifest = helper.prepare_original_input(root, spec, "zh-CN")
            (root / "findings.json").write_text(json.dumps({"findings": [{"original_input": spec}]}))
            doc = Document()
            helper.render_original_input(doc, root, manifest)
            doc.save(root / "report.docx")
            helper.validate_original_input(root, required=True)
            self.assertEqual(before, (root / spec["screenshot_path"]).read_bytes())
            import validate_recording_evidence as recording
            script = root / "run-demo.sh"
            script.write_text("# recording_checkpoint ZHULONG_RECORDING_STAGE_DIR\n")
            replay = {"replay": {"script_path": script.name, "script_sha256": helper.digest(script), "exit_code": 0}}
            with self.assertRaisesRegex(Exception, "ORIGINAL_INPUT_INVALID"):
                recording._validate_replay(root, replay)
            (root / "attachments/evidence").mkdir()
            (root / "attachments/evidence/replay-runtime-output.log").write_bytes(raw.read_bytes())
            recording._validate_replay(root, replay)
            raw.write_bytes(b"changed\n")
            with self.assertRaisesRegex(ValueError, "ORIGINAL_INPUT_INVALID"):
                helper.validate_original_input(root, required=True)

    def test_declared_input_cannot_disappear_and_docx_errors_are_stable(self):
        from docx import Document
        from PIL import Image
        helper = importlib.import_module("original_input_evidence")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "attachments").mkdir()
            (root / "attachments/input.txt").write_text("synthetic text\n")
            Image.new("RGB", (80, 40), "white").save(root / "attachments/input.png")
            spec = {"text_path": "attachments/input.txt", "screenshot_path": "attachments/input.png", "title": "PoC 原始输入", "description": "测试说明", "caption": "图 1：原始输入截图"}
            findings = root / "findings.json"
            findings.write_text(json.dumps({"findings": [{"original_input": spec}]}))
            with self.assertRaisesRegex(ValueError, "ORIGINAL_INPUT_INVALID"):
                helper.validate_original_input(root)
            manifest = helper.prepare_original_input(root, spec, "zh-CN")
            doc = Document()
            helper.render_original_input(doc, root, manifest)
            doc.save(root / "report.docx")
            changed = dict(spec, caption="图 2：被更改的声明")
            findings.write_text(json.dumps({"findings": [{"original_input": changed}]}))
            with self.assertRaisesRegex(ValueError, "ORIGINAL_INPUT_INVALID"):
                helper.validate_original_input(root)
            findings.write_text(json.dumps({"findings": [{"original_input": spec}]}))
            (root / "report.docx").write_bytes(b"broken document")
            with self.assertRaisesRegex(ValueError, "ORIGINAL_INPUT_INVALID"):
                helper.validate_original_input(root)

    def test_missing_input_is_not_required_delivery_and_english_caption_is_rejected(self):
        helper = importlib.import_module("original_input_evidence")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.assertIsNone(helper.validate_original_input(root))
            with self.assertRaisesRegex(ValueError, "required"):
                helper.validate_original_input(root, required=True)
            for caption in ("Figure: Full packet", "图 1：Review evidence"):
                spec = {"text_path": "attachments/input.txt", "screenshot_path": "attachments/input.png", "title": "PoC 原始输入", "description": "测试说明", "caption": caption}
                with self.assertRaisesRegex(ValueError, "Chinese"):
                    helper.prepare_original_input(root, spec, "zh-CN")

    def test_original_screenshot_rejects_unsafe_paths_and_links(self):
        from PIL import Image
        helper = importlib.import_module("original_input_evidence")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "attachments").mkdir()
            (root / "attachments/input.txt").write_text("synthetic input\n")
            source = root / "source.png"
            Image.new("RGB", (80, 40), "white").save(source)
            screenshot = root / "attachments/input.png"
            spec = {"text_path": "attachments/input.txt", "screenshot_path": "attachments/input.png", "title": "PoC 原始输入", "description": "测试说明", "caption": "图 1：原始输入截图"}
            for kind in ("symlink", "hardlink", "directory"):
                if kind == "symlink":
                    screenshot.symlink_to(source)
                elif kind == "hardlink":
                    os.link(source, screenshot)
                else:
                    screenshot.mkdir()
                with self.assertRaisesRegex(ValueError, "ORIGINAL_INPUT_INVALID"):
                    helper.prepare_original_input(root, spec, "zh-CN")
                screenshot.rmdir() if kind == "directory" else screenshot.unlink()
            with self.assertRaisesRegex(ValueError, "ORIGINAL_INPUT_INVALID"):
                helper.prepare_original_input(root, dict(spec, screenshot_path="attachments/../source.png"), "zh-CN")


class ProvisioningTests(unittest.TestCase):
    def test_language_reconstruction_failures_are_independent(self):
        import render_confirmed_vuln_docx as renderer
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "compose.yaml").write_text('{"services":{"app":{"build":"."}}}')
            (root / "Dockerfile").write_text("FROM scratch\n")
            artifact = {"output_name": "run-demo.sh", "generator": "reviewer-recording-shell"}
            finding = {"project_name": "demo", "reproduction": [{"commands": ["docker compose build app"]}], "bundle_root_artifacts": [artifact]}
            script = root / "run-demo.sh"
            for language, other, name in (("en-US", "zh-CN", "Path Traversal"), ("zh-CN", "en-US", "路径遍历")):
                with self.subTest(language=language):
                    finding.update(report_language=language, vulnerability_name=name, vuln_type=name)
                    with self.assertRaises(SystemExit):
                        renderer.build_generated_recording_shell(finding, other, {}, artifact)
                    shell = renderer.build_generated_recording_shell(finding, language, {}, artifact)
                    (root / "findings.json").write_text(json.dumps({"findings": [finding]}))
                    script.write_text(shell)
                    validator.validate_standalone_replay_inputs(root)
                    script.write_text(shell + "\n./provision.sh\n")
                    with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNSUPPORTED"):
                        validator.validate_standalone_replay_inputs(root)
                    finding["reproduction"][0]["commands"].append("./provision.sh")
                    (root / "findings.json").write_text(json.dumps({"findings": [finding]}))
                    script.write_text(renderer.build_generated_recording_shell(finding, language, {}, artifact))
                    with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNSUPPORTED"):
                        validator.validate_standalone_replay_inputs(root)
                    finding["reproduction"][0]["commands"].pop()
            script.write_text(shell)
            finding.update(vulnerability_name="", vuln_type="")
            for language in ("zh-CN", "en-US"):
                with self.assertRaises(SystemExit):
                    renderer.build_generated_recording_shell(finding, language, {}, artifact)
            (root / "findings.json").write_text(json.dumps({"findings": [finding]}))
            with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNSUPPORTED"):
                validator.validate_standalone_replay_inputs(root)

    def test_ascii_english_builder_and_strict_production_cli(self):
        import selftest_plugin as fixtures
        plugin = Path(__file__).resolve().parent.parent
        with tempfile.TemporaryDirectory(prefix="zhulong-language-cli-") as tmp:
            root = Path(tmp)
            fakebin = root / "bin"
            fakebin.mkdir()
            docker = fakebin / "docker"
            docker.write_text('#!/bin/sh\ncase "$*" in "compose version"|*" config") exit 0;; *) exit 99;; esac\n')
            docker.chmod(0o755)
            env = dict(os.environ, PATH=str(fakebin) + os.pathsep + os.environ["PATH"])
            repo, workspace = fixtures.new_build_wrapper_workspace(root, "english")
            slug = fixtures.build_wrapper_source_finding(plugin, repo, workspace, slug="demo-app_Path_Traversal_High_report")
            contract = fixtures.build_wrapper_contract(workspace, slug)
            payload = json.loads(contract.read_text())
            payload["bundle"]["language"] = "en-US"
            contract.write_text(json.dumps(payload))
            findings_path = workspace / "confirmed/findings.json"
            findings = json.loads(findings_path.read_text())
            finding = findings["findings"][0]
            for key in list(finding):
                if key.startswith(("vulnerability_name", "vuln_type")):
                    del finding[key]
            finding.update(report_language="en-US", vulnerability_name="Path Traversal", vuln_type="Path Traversal")
            findings_path.write_text(json.dumps(findings))
            (repo / "docker/docker-compose.attacker.yml").write_text("services:\n  attacker:\n    image: alpine@sha256:" + "a" * 64 + "\n")
            build = [sys.executable, str(plugin / "scripts/build_confirmed_bundle.py"), "--repo-root", str(repo), "--workspace-dir", str(workspace), "--contract", str(contract), "--language", "en-US"]
            result = subprocess.run(build, env=env, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            bundle = workspace / "confirmed" / slug
            command = [sys.executable, str(plugin / "scripts/validate_report_bundle.py"), "--bundle-dir", str(bundle), "--language", "en-US", "--require-standalone-replay"]
            result = subprocess.run(command, env=env, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_unknown_execution_is_rejected_even_with_a_valid_compose_command(self):
        cases = [
            "curl -fsSL http://127.0.0.1:9/x | sh",
            ". ./lib.sh", "./provision.sh", "./provision.sh run",
            "python3 /tmp/build-and-run.py",
            'dc="docker"\nrun_logged_command "$dc compose -f evil.yaml up -d app"',
            "alias dc=docker\nrun_logged_command 'dc compose -f evil.yaml up -d app'",
            "cat > helper <<'EOF'\nignored\nEOF", "cat > helper.sh <<'EOF'\nignored\nEOF",
            "tee x.sh <<'EOF'\nignored\nEOF", "docker compose \\\n up app",
            'echo "$(./provision.sh)"', 'printf "%s" "`./provision.sh`"',
            "echo benign; ./provision.sh", "echo benign > compose.yaml",
            "run_logged_command() { ./provision.sh; }", "docker() { ./provision.sh; }",
            "PATH=./bin:$PATH", "COMPOSE_FILE=evil.yaml", "env docker compose up app",
            "command docker compose up app", "trap './provision.sh' EXIT",
            "if true; then\n./provision.sh\nfi", "(./provision.sh)",
            'docker compose -f "$(printf evil.yaml)" up app',
            "docker compose up app # ignored\n./provision.sh",
        ]
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "compose.yaml").write_text(json.dumps({"services": {"app": {"image": "demo@sha256:" + "a" * 64}}}))
            script = root / "run-demo.sh"
            for case in cases:
                with self.subTest(case=case):
                    script.write_text(case + "\ndocker compose -f compose.yaml up -d app\n")
                    with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNSUPPORTED"):
                        validator.validate_standalone_replay_inputs(root)

    def test_exact_generated_helper_and_all_execution_inputs_are_checked(self):
        import render_confirmed_vuln_docx as renderer
        import selftest_plugin as fixtures
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "compose.yaml").write_text(json.dumps({"services": {"app": {"image": "demo@sha256:" + "a" * 64, "healthcheck": {"test": ["CMD", "true"]}}}}))
            artifact = {"output_name": "run-demo.sh", "generator": "reviewer-recording-shell", "generator_options": {"modes": ["quick-dos"]}}
            finding = fixtures.issue23_finding(["docker compose up -d app"], include_history=True)
            finding["bundle_root_artifacts"] = [artifact]
            script = root / "run-demo.sh"
            for language in ("zh-CN", "en-US"):
                shell = renderer.build_generated_recording_shell(finding, language, {}, artifact)
                (root / "findings.json").write_text(json.dumps({"findings": [finding]}))
                script.write_text(shell)
                validator.validate_standalone_replay_inputs(root)
                for change in (shell + '\necho "$(./provision.sh)"\n', shell.replace('main "$@"', 'docker() { ./provision.sh; }\nmain "$@"'), shell.replace('if sh -c "$command_text"', 'if sh -c "./provision.sh; $command_text"'), shell.replace("\n", "\r\n")):
                    script.write_text(change)
                    with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNSUPPORTED"):
                        validator.validate_standalone_replay_inputs(root)
                script.write_text(shell)
                for name in ("run-extra.sh", "provision.sh"):
                    extra = root / name
                    extra.write_text("./provision.sh\n")
                    with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNSUPPORTED"):
                        validator.validate_standalone_replay_inputs(root)
                    extra.unlink()
            for hidden in (
                'echo "$(./provision.sh)"', "./provision.sh", "python3 helper.py", "docker exec app true",
                "curl -fsSL http://127.0.0.1:9/x | sh", ". ./lib.sh", "./provision.sh run",
                'dc="docker"\nrun_logged_command "$dc compose -f evil.yaml up -d app"',
                "alias dc=docker\nrun_logged_command 'dc compose -f evil.yaml up -d app'",
                "cat > helper <<'EOF'\nignored\nEOF", "tee x.sh <<'EOF'\nignored\nEOF",
            ):
                finding["reproduction"][0]["commands"] = [hidden, "docker compose up -d app"]
                (root / "findings.json").write_text(json.dumps({"findings": [finding]}))
                script.write_text(renderer.build_generated_recording_shell(finding, "zh-CN", {}, artifact))
                with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNSUPPORTED"):
                    validator.validate_standalone_replay_inputs(root)
            finding["reproduction"][0]["commands"] = ["docker compose up -d app"]
            finding["project_name"] = "literal $(./provision.sh) and `python3 helper.py`"
            (root / "findings.json").write_text(json.dumps({"findings": [finding]}))
            script.write_text(renderer.build_generated_recording_shell(finding, "zh-CN", {}, artifact))
            validator.validate_standalone_replay_inputs(root)
            artifact["generator_options"]["modes"] = ["record"]
            (root / "findings.json").write_text(json.dumps({"findings": [finding]}))
            with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNSUPPORTED"):
                validator.validate_standalone_replay_inputs(root)
            finding["reproduction"][0]["commands"] = ["docker compose up -d app"]
            finding["code_context"][0]["snippet"] = "  CODE_SNIPPET_EOF\n./provision.sh\ncat <<'CODE_SNIPPET_EOF'"
            (root / "findings.json").write_text(json.dumps({"findings": [finding]}))
            script.write_text(renderer.build_generated_recording_shell(finding, "zh-CN", {}, artifact))
            with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNSUPPORTED"):
                validator.validate_standalone_replay_inputs(root)

    def test_minimal_literal_script_positive(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "compose.yaml").write_text(json.dumps({"services": {"app": {"build": "."}}}))
            (root / "Dockerfile").write_text("FROM scratch\n")
            (root / "run-demo.sh").write_text("#!/bin/sh\nset -eu\n# no host commands\ndocker compose -f 'compose.yaml' up -d app\n")
            validator.validate_standalone_replay_inputs(root)

    def test_builder_and_production_validator_cli(self):
        import selftest_plugin as fixtures
        from PIL import Image
        plugin = Path(__file__).resolve().parent.parent
        with tempfile.TemporaryDirectory(prefix="zhulong-input-cli-") as tmp:
            root = Path(tmp)
            fakebin = root / "bin"
            fakebin.mkdir()
            docker = fakebin / "docker"
            docker.write_text('#!/bin/sh\ncase "$*" in "compose version"|*" config") exit 0;; *) exit 99;; esac\n')
            docker.chmod(0o755)
            env = dict(os.environ, PATH=str(fakebin) + os.pathsep + os.environ["PATH"])
            repo, workspace = fixtures.new_build_wrapper_workspace(root, "cli")
            slug = fixtures.build_wrapper_source_finding(plugin, repo, workspace)
            contract = fixtures.build_wrapper_contract(workspace, slug)
            findings_path = workspace / "confirmed/findings.json"
            findings = json.loads(findings_path.read_text())
            finding = findings["findings"][0]
            (repo / "evidence/input.txt").write_bytes(b"GET /synthetic HTTP/1.1\r\nHost: example.invalid\r\n\r\n")
            Image.new("RGB", (160, 80), "white").save(repo / "evidence/input.png")
            finding["attachments"].extend([{"path": "evidence/input.txt", "purpose": "原始输入测试文本"}, {"path": "evidence/input.png", "purpose": "独立合成测试图片"}])
            finding["original_input"] = {"text_path": "evidence/input.txt", "screenshot_path": "evidence/input.png", "title": "PoC 原始输入", "description": "此图片仅验证交付格式，不是真实漏洞截图。", "caption": "图 1：原始输入测试图片"}
            findings_path.write_text(json.dumps(findings, ensure_ascii=False))
            # Pin the delivered attacker service. No image is pulled or executed.
            (repo / "docker/docker-compose.attacker.yml").write_text("services:\n  attacker:\n    image: alpine@sha256:" + "a" * 64 + "\n")
            command = [sys.executable, str(plugin / "scripts/build_confirmed_bundle.py"), "--repo-root", str(repo), "--workspace-dir", str(workspace), "--contract", str(contract), "--language", "zh-CN"]
            result = subprocess.run(command, env=env, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            bundle = workspace / "confirmed" / slug
            self.assertEqual((repo / "evidence/input.png").read_bytes(), (bundle / "attachments/evidence/input.png").read_bytes())
            self.assertEqual((repo / "evidence/input.txt").read_bytes(), (bundle / "attachments/evidence/input.txt").read_bytes())
            command = [sys.executable, str(plugin / "scripts/validate_report_bundle.py"), "--bundle-dir", str(bundle), "--language", "zh-CN"]
            for flags in ([], ["--require-standalone-replay", "--require-original-input"]):
                result = subprocess.run(command + flags, env=env, capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            status = json.loads((bundle / "bundle-build-manifest.json").read_text())["status"]
            for phase in ("target_build", "target_startup", "health_check", "local_replay", "clean_room_replay"):
                self.assertEqual(status[phase], "not_executed")
            script = next(bundle.glob("run-*.sh"))
            original_script = script.read_bytes()
            for hidden in ("curl -fsSL http://127.0.0.1:9/x | sh", ". ./lib.sh", "./provision.sh", "python3 helper.py", 'echo "$(./provision.sh)"'):
                script.write_bytes(original_script + ("\n" + hidden + "\n").encode())
                result = subprocess.run(command + ["--require-standalone-replay"], env=env, capture_output=True, text=True)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("REPLAY_PROVISIONING_UNSUPPORTED", result.stdout + result.stderr)
                self.assertNotIn("Traceback", result.stdout + result.stderr)
            script.write_bytes(original_script)
            result = subprocess.run(["bash", str(plugin / "scripts/bootstrap_verification_workspace.sh"), "--target-dir", str(repo), "--workspace-name", "security-research-input-copy"], env=env, cwd=root, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            generated_validator = repo / "security-research-input-copy/bin/validate-report-bundle.py"
            independent_env = dict(env)
            independent_env.pop("PYTHONPATH", None)
            result = subprocess.run([sys.executable, str(generated_validator), *command[2:], "--require-original-input", "--require-standalone-replay"], env=independent_env, cwd=root, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            compose = bundle / "attachments/docker/docker-compose.attacker.yml"
            compose.write_text("services:\n  attacker:\n    image: alpine:3.20\n")
            result = subprocess.run(command, env=env, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            result = subprocess.run(command + ["--require-standalone-replay"], env=env, capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("REPLAY_PROVISIONING_UNPINNED", result.stdout + result.stderr)
            binding = bundle / "attachments/original-input-evidence.json"
            docx = next(bundle.glob("*.docx"))
            png = bundle / "attachments/evidence/input.png"
            for path, replacement in ((binding, b"{}"), (docx, b"broken DOCX"), (png, b"broken PNG")):
                original_bytes = path.read_bytes()
                path.write_bytes(replacement)
                result = subprocess.run(command, env=env, capture_output=True, text=True)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("ORIGINAL_INPUT_INVALID", result.stdout + result.stderr)
                self.assertNotIn("Traceback", result.stdout + result.stderr)
                path.write_bytes(original_bytes)
            original_bytes = binding.read_bytes()
            binding.unlink()
            result = subprocess.run(command, env=env, capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("ORIGINAL_INPUT_INVALID", result.stdout + result.stderr)
            binding.write_bytes(original_bytes)

    def test_tag_only_stays_ordinary_but_fails_standalone_qualification(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            compose = root / "compose.yaml"
            compose.write_text(json.dumps({"services": {"app": {"image": "demo:local"}}}))
            validator.validate_compose_static_paths(compose, root)
            with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNPINNED"):
                validator.validate_replay_compose_inputs(root, ["up", "app"], require_standalone=True)

    def test_qualification_checks_dependencies_and_merged_build(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            base = root / "compose.yaml"
            base.write_text(json.dumps({"services": {"app": {"image": "demo:local", "depends_on": ["db"]}, "db": {"image": "db@sha256:" + "a" * 64}}}))
            override = root / "extra.yaml"
            override.write_text(json.dumps({"services": {"app": {"build": "."}}}))
            (root / "Dockerfile").write_text("FROM scratch\n")
            validator.validate_replay_compose_inputs(root, ["-f", "compose.yaml", "-f", "extra.yaml", "up"], require_standalone=True)
            base.write_text(json.dumps({"services": {"app": {"build": ".", "depends_on": ["db"]}, "db": {"image": "db:local"}}}))
            with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNPINNED"):
                validator.validate_replay_compose_inputs(root, ["up", "app"], require_standalone=True)

    def test_qualification_rejects_external_resources_and_unknown_shell(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "compose.yaml").write_text(json.dumps({"services": {"app": {"image": "demo@sha256:" + "a" * 64}}, "networks": {"default": {"external": True}}}))
            with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNSUPPORTED"):
                validator.validate_replay_compose_inputs(root, ["up"], require_standalone=True)
            script = root / "run-demo.sh"
            script.write_text("docker run demo:local\n")
            with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNSUPPORTED"):
                validator.validate_standalone_replay_inputs(root)

    def test_shell_tail_cannot_hide_a_local_image_dependency(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "compose.yaml").write_text(json.dumps({"services": {"app": {"image": "demo@sha256:" + "a" * 64}}}))
            (root / "run-demo.sh").write_text("docker compose up app; docker run hidden:local\n")
            with self.assertRaisesRegex(SystemExit, "REPLAY_PROVISIONING_UNSUPPORTED"):
                validator.validate_standalone_replay_inputs(root)


if __name__ == "__main__":
    unittest.main()
