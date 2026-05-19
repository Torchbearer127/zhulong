#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Iterable

try:
    from docx import Document
except ModuleNotFoundError as exc:
    raise SystemExit(
        "python-docx is required. Install it with `python3 -m pip install python-docx` and rerun."
    ) from exc


REQUIRED_DOCX_MEMBERS = {"[Content_Types].xml", "word/document.xml", "_rels/.rels"}
ZH_HEADINGS = {"漏洞描述", "影响版本", "漏洞危险性评估", "漏洞分析", "漏洞复现", "最终判定："}
EN_HEADINGS = {
    "Vulnerability Description",
    "Affected Versions",
    "Risk Assessment",
    "Vulnerability Analysis",
    "Reproduction",
    "Final Verdict:",
}
ABSOLUTE_PATH_PATTERNS = [
    re.compile(r"/Users/[^/\s]+"),
    re.compile(r"/home/[^/\s]+"),
    re.compile(r"[A-Za-z]:\\\\"),
    re.compile(r"file://"),
]
GENERIC_DOCX_FILENAMES = {"report.docx", "vulnerability-report.docx", "漏洞报告.docx"}
GENERIC_NOTE_FILENAMES = {"attachments.md", "attachment-index.md", "附件目录说明.md"}
GENERIC_SUPPLEMENT_FILENAMES = {"reproduction.md", "reproduction_note.md", "补充复现说明.md"}
PLACEHOLDER_REPORT_TEXT = {
    "zh-CN": [
        "最终判定待补充",
        "漏洞分析待补充",
        "评估依据待补充",
        "复现步骤待补充",
        "关键代码位置待补充",
        "影响信息待补充",
        "CVSS 信息待补充",
    ],
    "en-US": [
        "Final verdict pending",
        "Vulnerability analysis is not provided yet",
        "Assessment rationale is not provided yet",
        "Reproduction steps are not provided yet",
        "Key code path is not provided yet",
        "Impact details are not provided yet",
        "CVSS details are not provided yet",
    ],
}
HOST_FALLBACK_PATTERNS = [
    re.compile(r"\[record\|quick\]\s+\[docker\|host\]"),
    re.compile(r"\brecord host\b", re.IGNORECASE),
    re.compile(r"\bquick host\b", re.IGNORECASE),
    re.compile(r"runtime=host", re.IGNORECASE),
    re.compile(r"\[docker\|host\]"),
]
SCRIPT_PAUSE_PATTERNS = [
    re.compile(r"\bpause_step\s*\("),
    re.compile(r"\bsleep\s+['\"]?\$?[A-Za-z0-9_{-]+"),
]
SCRIPT_RUNTIME_PAUSE_PATTERNS = [
    re.compile(r"\bpause_step\b"),
    re.compile(r"\bsleep\s+['\"]?\$?[A-Za-z0-9_{-]+"),
]
SCRIPT_COLOR_PATTERNS = [
    re.compile(r"\\033\[[0-9;]+m"),
    re.compile(r"\$'\\033\[[0-9;]+m"),
    re.compile(r"\btput\s+setaf\b"),
    re.compile(r"\bC_(?:RED|GREEN|YELLOW|BLUE|MAGENTA|CYAN|RESET|BOLD)\b"),
]
SCRIPT_STEP_PATTERNS = [
    re.compile(r"Step\s+\d+/\d+"),
    re.compile(r"步骤\s*\d+/\d+"),
    re.compile(r"\bannounce_step\b"),
    re.compile(r"\bprint_banner\b"),
]
ZH_REVIEW_MARKERS = ["模式：", "运行环境：", "关键代码位置", "执行命令：", "证据汇总"]
EN_REVIEW_MARKERS = ["Mode:", "Runtime:", "Key code location:", "Run command:", "Evidence Summary"]
ZH_SUPPLEMENT_MARKERS = ["补充复现说明", "建议的最短复现路径", "关键成功证据"]
EN_SUPPLEMENT_MARKERS = ["Reproduction Supplement", "Recommended Shortest Reproduction Path", "Key Success Evidence"]
FORBIDDEN_BUNDLE_DIRS = {".omc", ".git", ".hg", ".svn", "__pycache__", "node_modules", ".venv", "venv", "evidence"}
ZH_ANALYSIS_MARKERS = ["位置：", "入口/可控输入：", "危险函数/危险操作：", "触发路径：", "根因：", "现有校验为何失效："]
EN_ANALYSIS_MARKERS = [
    "Location:",
    "Entry / controllable input:",
    "Dangerous operation:",
    "Trigger path:",
    "Root cause:",
    "Why existing checks fail:",
]
ZH_REPRODUCTION_MARKERS = ["环境准备", "执行", "预期", "实际", "结果证据"]
EN_REPRODUCTION_MARKERS = ["Environment", "Run command", "Expected", "Observed", "Evidence"]
ZH_QUALITY_LABELS = {
    "attacker": ["攻击者条件", "攻击前提", "攻击者前提"],
    "server": ["服务端条件", "服务端前提", "服务端配置条件"],
    "impact": ["安全影响", "安全影响说明"],
}
EN_QUALITY_LABELS = {
    "attacker": ["Attacker Condition", "Attack Preconditions", "Attacker Preconditions"],
    "server": ["Server Condition", "Server Preconditions", "Server Configuration Condition"],
    "impact": ["Security Impact", "Confirmed Security Impact"],
}
ZH_LEGACY_QUALITY_MARKERS = {
    "attacker": ["入口/可控输入："],
    "server": ["服务端条件：", "服务端前提：", "默认配置", "默认 IMPORT_IP_DENY_LIST", "服务端配置"],
    "impact": ["实际影响/边界：", "评估依据：", "结果证据：", "安全影响："],
}
EN_LEGACY_QUALITY_MARKERS = {
    "attacker": ["Entry / controllable input:"],
    "server": ["Server condition:", "Server precondition:", "default configuration", "server configuration"],
    "impact": ["Practical impact / boundary:", "Assessment rationale:", "Evidence:", "Security impact:"],
}
ZH_QUALITY_PLACEHOLDERS = {
    "待补充",
    "暂无",
    "无",
    "未知",
    "不详",
    "请补充",
    "占位",
    "todo",
    "tbd",
    "n/a",
}
EN_QUALITY_PLACEHOLDERS = {
    "pending",
    "todo",
    "tbd",
    "n/a",
    "na",
    "none",
    "unknown",
    "placeholder",
    "fillin",
    "notprovided",
}
ZH_IMPACT_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"机密性",
        r"数据泄露|信息泄露|敏感(?:信息|文件|数据)?(?:读取|泄露|暴露)?",
        r"读取|窃取|泄露|暴露|内网|SSRF",
        r"完整性|未授权|越权|篡改|修改|写入|伪造|权限",
        r"可用性|拒绝服务|DoS|崩溃|资源耗尽",
    )
]
EN_IMPACT_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"confidentiality|data exposure|information disclosure",
        r"sensitive (?:file|data|information)|read|exfiltrat|leak|disclos",
        r"integrity|unauthori[sz]ed|authorization|modify|write|tamper|forge|privilege",
        r"availability|denial of service|DoS|crash|resource exhaustion",
        r"SSRF|internal service|internal network",
    )
]
ALLOWED_VERIFICATION_STATUSES = {
    "confirmed_in_docker",
    "high_confidence_unverified_due_to_sandbox_limitation",
}
PLACEHOLDER_VERIFICATION_VALUES = {
    "project-specific Docker image or Docker Compose service",
}
REQUIRED_VERIFICATION_FIELDS = {
    "schema_version",
    "finding_slug",
    "verification_status",
    "docker_required",
    "docker_image",
    "docker_command",
    "poc_path",
    "expected_observation",
    "observed_observation",
    "oracle_token",
    "evidence_files",
    "severity_escalation_attempted",
    "severity_escalation_result",
}
BROAD_AFFECTED_VERSION_PATTERNS = [
    re.compile(r"\ball\s+versions\b", re.IGNORECASE),
    re.compile(r"\bpotentially\s+all\s+versions\b", re.IGNORECASE),
    re.compile(r"所有版本"),
    re.compile(r"可能影响所有版本"),
]
RUNTIME_MISMATCH_TEXT_PATTERNS = [
    re.compile(r"source/runtime mismatch", re.IGNORECASE),
    re.compile(r"runtime/source mismatch", re.IGNORECASE),
    re.compile(r"source[- ]runtime mismatch", re.IGNORECASE),
    re.compile(r"源码/运行时不匹配"),
    re.compile(r"源代码/运行时不匹配"),
    re.compile(r"运行时.*不匹配"),
]
UNAUTHENTICATED_TITLE_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"\bunauthenticated\b",
        r"\bno[- ]auth\b",
        r"\bpre[- ]auth(?:enticated|entication)?\b",
        r"未认证",
        r"未授权",
        r"无需认证",
        r"认证前",
    )
]
AUTHENTICATED_FLOW_MARKERS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"/auth/login",
        r"Authorization:\s*Bearer",
        r"\baccess_token\b",
        r"\brefresh_token\b",
    )
]
RECORDING_ZERO_STEP_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"\[0/\d+\]",
        r"\b0/\d+\b",
        r"\bStep\s+0/\d+\b",
        r"步骤\s*0/\d+",
    )
]
SHELL_CONFIRMATION_PHRASE_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"\bVULNERABILITY CONFIRMED\b",
        r"\bATTACK SUCCESS\b",
        r"漏洞已确认",
        r"攻击成功",
    )
]
SHELL_CONFIRMATION_PRINT_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"\b(?:echo|printf|highlight_success|highlight_danger|print_banner)\b.*(?:VULNERABILITY CONFIRMED|ATTACK SUCCESS)\b",
        r"\b(?:echo|printf|highlight_success|highlight_danger|print_banner)\b.*(?:漏洞已确认|攻击成功)",
    )
]
SHELL_SUCCESS_ORACLE_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"\bjq\s+-e\b",
        r"\bgrep\s+-q\b",
        r"\btest\s+",
        r"\[\[",
        r"^\s*case\s+",
        r"^\s*if\b",
        r";\s*then\b",
        r"\bHTTP[_ -]?(?:status|code)?\b",
        r"\bstatus_code\b",
        r"\bresponse[_-]?(?:body|json|content)?\b",
        r"\bjson\b",
        r"\b(?:require|assert|check|verify)[A-Za-z0-9_ -]*(?:success|oracle|status|response|json|ok)\b",
        r"\bexit\s+1\b",
    )
]
SHELL_FAIL_OPEN_ORACLE_COMMAND_PATTERN = re.compile(
    r"\b(?:grep|jq|curl|docker\s+logs|test)\b|\[\[|\[\s",
    re.IGNORECASE,
)
SHELL_FAIL_OPEN_SOFTENER_PATTERN = re.compile(r"\|\|\s*(?:echo|printf|true|:)\b", re.IGNORECASE)
SHELL_FAIL_CLOSED_SOFTENER_PATTERN = re.compile(r"\|\|.*\b(?:exit|return)\s+[1-9][0-9]*\b", re.IGNORECASE)
COMPOSE_FILE_PATTERN = re.compile(r"^(?:docker-compose.*|compose.*)\.ya?ml$", re.IGNORECASE)
TECHNICAL_TEXT_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"https?://",
        r"\b(?:curl|docker|python3?|jq|grep|sed|awk|npm|yarn|pnpm|uv|bash|sh)\b",
        r"\b(?:GET|POST|PUT|PATCH|DELETE|HEAD|OPTIONS)\s+/",
        r"\b(?:Authorization|Content-Type|Accept|Cookie|User-Agent):",
        r"\bCVSS:\d\.\d/",
        r"\b[A-Za-z_][A-Za-z0-9_]*=",
        r"[{}[\]<>`$]",
        r"(?:^|\s)/(?:[A-Za-z0-9_.-]+/)+",
        r"\b(?:VULNERABILITY CONFIRMED|ATTACK SUCCESS)\b",
        r"漏洞已确认|攻击成功",
    )
]
STRUCTURED_TARGET_FIELDS = ("attack_target", "ssrf_target", "target_url")
ZH_REAL_WORLD_HEADINGS = {
    "实际场景中的危害与利用方式",
    "实际利用场景",
    "真实利用链",
    "攻击路径与影响外显",
}
EN_REAL_WORLD_HEADINGS = {
    "Real-World Exploitability",
    "Practical Exploitation Path",
    "Real-World Impact",
    "Attack Path and Observable Impact",
    "Practical Impact and Exploitation Path",
}
ZH_EXPLOITABILITY_MARKERS = {
    "attacker": ("攻击者", "用户", "可控", "控制", "污染", "提交", "访问", "认证", "权限", "输入", "参数"),
    "server": ("服务端", "服务器", "后端", "运行", "启用", "配置", "默认", "部署", "依赖", "日志", "错误", "调试"),
    "observable": ("日志", "错误", "响应", "返回", "输出", "回显", "可见", "外显", "存储", "记录", "回调", "流量", "观察", "证据"),
    "boundary": ("已验证", "验证", "Docker", "PoC", "证明", "确认", "边界", "限定", "不声称", "未验证", "不扩大", "仅"),
}
EN_EXPLOITABILITY_MARKERS = {
    "attacker": ("attacker", "user", "control", "influence", "poison", "submit", "access", "authenticated", "privilege", "input"),
    "server": ("server", "runtime", "configuration", "default", "deploy", "dependency", "feature", "debug", "logging", "enabled"),
    "observable": ("log", "error", "response", "output", "observable", "visible", "stored", "record", "callback", "traffic", "evidence"),
    "boundary": ("verified", "docker", "poc", "confirmed", "proven", "boundary", "limited", "not claim", "not verified", "stronger"),
}
EXPLOITABILITY_DISPLAY = {
    "zh-CN": {
        "attacker": "攻击者条件/可控输入",
        "server": "服务端可达条件",
        "observable": "影响外显或安全相关可见通道",
        "boundary": "已验证影响边界/未声称影响",
    },
    "en-US": {
        "attacker": "attacker condition / controllable input",
        "server": "server-side reachability condition",
        "observable": "observable or security-relevant impact channel",
        "boundary": "verified impact boundary / non-claimed impact",
    },
}
STRONG_ATTACKER_CONTROL_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"arbitrary code execution",
        r"malicious\s+JS",
        r"恶意\s*JS",
        r"攻击者已能执行",
        r"直接执行(?:恶意|攻击者|JS|Python|PHP|可执行|本地源文件)",
        r"写入可执行文件",
        r"attacker controls? a local source file",
        r"attacker-controlled local source file",
    )
]
PACKAGE_MANAGER_INSTALL_PATTERN = re.compile(r"\b(?:npm|yarn|pnpm)\s+(?:install|i|ci|add)\b", re.IGNORECASE)
INSTALL_LIFECYCLE_SAFE_PATTERN = re.compile(
    r"--ignore-scripts|--offline|--prefer-offline|file:|\.tgz\b|attachments/|vendor/",
    re.IGNORECASE,
)
SCRIPT_PARENT_TRAVERSAL_PATTERN = re.compile(
    r"(?P<token>(?:\$\{?(?:SCRIPT_DIR|BUNDLE_ROOT|ATTACH_DIR|PWD)\}?/)?(?:\.\./)+(?:[A-Za-z0-9_.${}/-]+)?)"
)
POC_LABEL_PATTERN = re.compile(r"\bP(?:o)?C[-\s]*(\d+)\b", re.IGNORECASE)
VIDEO_SUFFIXES = {".mp4", ".mov", ".webm", ".mkv"}
TEXT_EVIDENCE_SUFFIXES = {
    ".bash",
    ".cjs",
    ".conf",
    ".env",
    ".ini",
    ".js",
    ".json",
    ".log",
    ".mjs",
    ".md",
    ".php",
    ".py",
    ".rb",
    ".sh",
    ".txt",
    ".yaml",
    ".yml",
}
REVIEWER_EVIDENCE_ADDENDUM = "reviewer-evidence-and-impact.md"
REVIEWER_EVIDENCE_INDEX = Path("attachments/reviewer-evidence-index.json")
REVIEWER_INDEX_ARTIFACT_KEYS = {
    "artifact",
    "artifacts",
    "artifact_file",
    "artifact_files",
    "artifact_path",
    "artifact_paths",
    "build_context",
    "evidence_file",
    "evidence_files",
    "evidence_output",
    "evidence_outputs",
    "file",
    "files",
    "fixture_file",
    "fixture_files",
    "media",
    "poc_file",
    "poc_files",
    "source_evidence",
    "source_file",
    "source_files",
    "vendor_fixture",
    "video_evidence",
    "video_file",
    "video_files",
}
REVIEWER_INDEX_ORACLE_KEYS = {
    "oracle_token",
    "oracle_tokens",
    "success_oracle",
    "success_oracles",
    "success_token",
    "success_tokens",
}
REVIEWER_COMMAND_KEYS = (
    "bundle_root_command",
    "replay_command",
    "reviewer_command",
    "shortest_replay_command",
)
REVIEWER_ADDENDUM_MARKERS = {
    "boundary": (
        "attacker",
        "capability",
        "boundary",
        "precondition",
        "condition",
        "not claim",
        "not verified",
        "攻击者",
        "前提",
        "条件",
        "边界",
        "不声称",
        "未验证",
        "不主张",
    ),
    "impact": ("impact", "harm", "verified", "影响", "危害", "已验证", "实测"),
    "oracle": ("oracle", "success token", "evidence", "成功判据", "证据", "token"),
    "replay": ("./run-", "docker", "replay command", "复现命令", "最短复现", "审核方最短"),
}
REVIEWER_PLACEHOLDER_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"\bTODO\b",
        r"\bTBD\b",
        r"\bplaceholder\b",
        r"待补充",
        r"请补充",
        r"占位",
        r"暂无",
    )
]
SUBMITTER_LOCAL_IMAGE_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"\bselftest:local(?:[-_.A-Za-z0-9]*)?\b",
        r"\b[A-Za-z0-9_.-]*selftest[A-Za-z0-9_.-]*:local(?:[-_.A-Za-z0-9]*)?\b",
        r"\b[A-Za-z0-9_.-]+:local(?:[-_.A-Za-z0-9]*)\b",
        r"\b[A-Za-z0-9_.-]+:latest-local(?:[-_.A-Za-z0-9]*)?\b",
    )
]
REMOTE_LATEST_IMAGE_PATTERN = re.compile(
    r"\b(?:docker\s+(?:run|pull|build|compose)|image:)\b[^\n`'\"]*:[Ll]atest\b",
    re.IGNORECASE,
)
FIXTURE_HINT_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"\bminimal\s+fixture\b",
        r"\bfixture[- ]based\b",
        r"\bvendored?\b",
        r"\bvendor/",
        r"attachments/vendor",
        r"attachments/fixture",
        r"uses_bundle_local_fixture",
        r"uses_bundle_local_vendor",
        r"最小\s*fixture",
        r"内置\s*fixture",
    )
]
LIBRARY_HINT_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"\blibrary vulnerability\b",
        r"\blibrary/package\b",
        r"\bpublic API\b",
        r"\bconsumer application\b",
        r"\bpackage vulnerability\b",
        r"\bprototype pollution\b",
        r"\bObject\.prototype\b",
        r"\bQ\.(?:set|get|post|invoke)\b",
        r"库漏洞",
        r"库本身",
        r"上层应用",
        r"消费方",
        r"原型链污染",
        r"全局\s*Object\.prototype",
    )
]
NONCLAIM_MARKERS = (
    "not claim",
    "not claimed",
    "not verified",
    "does not claim",
    "no direct",
    "未验证",
    "不声称",
    "不主张",
    "不作为",
    "不应",
    "不是",
    "没有在本包中验证",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Validate a confirmed vulnerability report bundle for structural and content stability."
    )
    parser.add_argument("--bundle-dir", required=True, help="Per-vulnerability bundle directory under confirmed/.")
    parser.add_argument(
        "--language",
        choices=["zh-CN", "en-US", "auto"],
        default="auto",
        help="Expected output language. Defaults to auto-detect from filenames/content.",
    )
    parser.add_argument(
        "--with-libreoffice",
        action="store_true",
        help="If LibreOffice is installed, also perform a headless DOCX->PDF conversion test.",
    )
    parser.add_argument(
        "--with-markitdown",
        action="store_true",
        help="If MarkItDown is installed, also perform a DOCX->Markdown extraction smoke test.",
    )
    parser.add_argument(
        "--write-audit-event",
        action="store_true",
        help=(
            "After successful validation, append a bundle_validated event to the workspace audit log. "
            "This is opt-in so standalone/offline bundle validation keeps its historical behavior."
        ),
    )
    return parser.parse_args()


def fail(message: str) -> None:
    raise SystemExit(f"VALIDATION FAILED: {message}")


def warn(message: str) -> None:
    print(f"WARN: {message}")


def find_audit_event_writer() -> Path | None:
    script_dir = Path(__file__).resolve().parent
    candidates = [
        script_dir / "write_audit_event.py",
        script_dir / "write-audit-event.py",
        script_dir.parent / "bin" / "write-audit-event.py",
    ]
    return next((path for path in candidates if path.exists()), None)


def write_bundle_validated_event(
    workspace_dir: Path,
    bundle_dir: Path,
    language: str,
    docx_path: Path,
    note_path: Path,
    supplement_path: Path,
    verification_evidence: dict[str, object],
) -> None:
    writer = find_audit_event_writer()
    if writer is None:
        fail("cannot write audit event because write_audit_event.py was not found next to the validator")

    details = {
        "bundle": f"confirmed/{bundle_dir.name}",
        "language": language,
        "verification_status": str(verification_evidence.get("verification_status") or ""),
        "finding_slug": str(verification_evidence.get("finding_slug") or ""),
        "docx": docx_path.name,
        "attachment_note": note_path.name,
        "reproduction_supplement": supplement_path.name,
    }
    command = [
        sys.executable,
        str(writer),
        "--workspace-dir",
        str(workspace_dir),
        "--event",
        "bundle_validated",
        "--stage",
        "reporting",
        "--status",
        "running",
        "--event-status",
        "ok",
        "--message",
        f"Confirmed bundle validated: {bundle_dir.name}",
        "--details-json",
        json.dumps(details, ensure_ascii=False, sort_keys=True),
    ]
    proc = subprocess.run(command, capture_output=True, text=True)
    if proc.returncode != 0:
        output = ((proc.stdout or "") + (proc.stderr or "")).strip()
        fail(f"failed to write bundle_validated audit event: {output}")


def docx_text(docx_path: Path) -> list[str]:
    doc = Document(docx_path)
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


def find_single(paths: list[Path], label: str) -> Path:
    if not paths:
        fail(f"missing {label}")
    if len(paths) > 1:
        fail(f"expected exactly one {label}, found {len(paths)}")
    return paths[0]


def detect_language(docx_path: Path, lines: Iterable[str], note_path: Path | None) -> str:
    if "漏洞报告" in docx_path.name:
        return "zh-CN"
    if docx_path.name.endswith("_report.docx"):
        return "en-US"
    content = "\n".join(lines)
    if "漏洞描述" in content or (note_path and "附件目录说明" in note_path.name):
        return "zh-CN"
    return "en-US"


def validate_docx_container(docx_path: Path) -> None:
    if not zipfile.is_zipfile(docx_path):
        fail(f"{docx_path.name} is not a valid OOXML/ZIP container")
    with zipfile.ZipFile(docx_path) as zf:
        names = set(zf.namelist())
    missing = REQUIRED_DOCX_MEMBERS - names
    if missing:
        fail(f"{docx_path.name} is missing OOXML members: {sorted(missing)}")


def validate_required_headings(lines: list[str], language: str) -> None:
    joined = "\n".join(lines)
    required = ZH_HEADINGS if language == "zh-CN" else EN_HEADINGS
    missing = [heading for heading in required if heading not in joined]
    if missing:
        fail(f"missing required headings for {language}: {missing}")


def section_text(lines: list[str], heading: str, next_headings: set[str]) -> str:
    start = None
    for index, line in enumerate(lines):
        if line.strip() == heading:
            start = index + 1
            break
    if start is None:
        return ""
    end = len(lines)
    for index in range(start, len(lines)):
        if lines[index].strip() in next_headings:
            end = index
            break
    return "\n".join(lines[start:end]).strip()


def validate_report_depth(lines: list[str], language: str) -> None:
    if language == "zh-CN":
        headings = ZH_HEADINGS | {"关键代码上下文", "验证环境关键文件"}
        analysis_heading = "漏洞分析"
        reproduction_heading = "漏洞复现"
        analysis_markers = ZH_ANALYSIS_MARKERS
        reproduction_markers = ZH_REPRODUCTION_MARKERS
    else:
        headings = EN_HEADINGS | {"Key Code Context", "Verification Environment Key Files"}
        analysis_heading = "Vulnerability Analysis"
        reproduction_heading = "Reproduction"
        analysis_markers = EN_ANALYSIS_MARKERS
        reproduction_markers = EN_REPRODUCTION_MARKERS

    analysis = section_text(lines, analysis_heading, headings)
    reproduction = section_text(lines, reproduction_heading, headings)
    analysis_hits = sum(1 for marker in analysis_markers if marker in analysis)
    reproduction_hits = sum(1 for marker in reproduction_markers if marker in reproduction)

    if len(analysis) < 220 or analysis_hits < 4:
        fail(
            "report Vulnerability Analysis section is too thin; it must explain location, controllable input, "
            "dangerous operation, trigger path, root cause, and why existing checks fail"
        )
    if len(reproduction) < 260 or reproduction_hits < 4:
        fail(
            "report Reproduction section is too thin; it must include setup, exact execution commands, "
            "expected result, observed result, and explicit success evidence"
        )


def validate_workspace_metadata(bundle_dir: Path) -> Path:
    confirmed_dir = bundle_dir.parent
    if confirmed_dir.name != "confirmed":
        fail(f"bundle must live under a confirmed/ directory, got: {bundle_dir}")
    workspace_dir = confirmed_dir.parent
    config_path = workspace_dir / "asr-config.json"
    if not config_path.exists():
        fail(f"workspace config missing: {config_path}")
    try:
        data = json.loads(config_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        fail(f"workspace config is unreadable: {config_path} ({exc})")
    if not isinstance(data, dict):
        fail("workspace config must be a JSON object")
    workspace_root = str(data.get("workspace_root", "")).strip()
    workspace_created_at = str(data.get("workspace_created_at", "")).strip()
    confirmed_output_dir = str(data.get("confirmed_output_dir", "")).strip()
    if workspace_dir.name == "security-research":
        fail("legacy bare repo/security-research workspace is not allowed for final deliverables")
    if workspace_root != workspace_dir.name:
        fail(
            f"workspace_root mismatch in asr-config.json: expected {workspace_dir.name}, got {workspace_root or '<missing>'}"
        )
    if not workspace_created_at:
        fail("workspace_created_at is missing in asr-config.json; this workspace looks legacy or inconsistent")
    if confirmed_output_dir != f"{workspace_dir.name}/confirmed":
        fail(
            "confirmed_output_dir mismatch in asr-config.json: "
            f"expected {workspace_dir.name}/confirmed, got {confirmed_output_dir or '<missing>'}"
        )
    return workspace_dir


def validate_output_filenames(docx_path: Path, note_path: Path, language: str) -> None:
    if docx_path.name in GENERIC_DOCX_FILENAMES:
        fail(f"generic docx filename is not allowed: {docx_path.name}")
    if note_path.name in GENERIC_NOTE_FILENAMES:
        fail(f"generic attachment note filename is not allowed: {note_path.name}")
    if language == "zh-CN" and "漏洞报告" not in docx_path.stem:
        fail(f"Chinese report filename must be finding-specific and include 漏洞报告: {docx_path.name}")
    if language == "en-US" and not docx_path.stem.endswith("_report"):
        fail(f"English report filename must be finding-specific and end with _report: {docx_path.name}")
    expected_note = f"{docx_path.stem}{'_attachment_index.md' if language == 'en-US' else '_附件目录说明.md'}"
    if note_path.name != expected_note:
        fail(f"attachment note filename must match the report stem: expected {expected_note}, got {note_path.name}")


def expected_supplement_filename(docx_path: Path, language: str) -> str:
    return f"{docx_path.stem}{'_reproduction_note.md' if language == 'en-US' else '_补充复现说明.md'}"


def normalize_token(text: str) -> str:
    return "".join(ch.lower() for ch in str(text) if ch.isalnum() or "\u4e00" <= ch <= "\u9fff")


def strip_quality_prefix(line: str) -> str:
    return re.sub(r"^[\s#>*\-•\d.、)）]+", "", str(line).strip()).strip()


def quality_labels(language: str) -> dict[str, list[str]]:
    return ZH_QUALITY_LABELS if language == "zh-CN" else EN_QUALITY_LABELS


def legacy_quality_markers(language: str) -> dict[str, list[str]]:
    return ZH_LEGACY_QUALITY_MARKERS if language == "zh-CN" else EN_LEGACY_QUALITY_MARKERS


def quality_display_name(language: str, kind: str) -> str:
    if language == "zh-CN":
        return {"attacker": "攻击者条件", "server": "服务端条件", "impact": "安全影响"}[kind]
    return {"attacker": "Attacker Condition", "server": "Server Condition", "impact": "Security Impact"}[kind]


def split_quality_label_line(line: str, labels: list[str]) -> tuple[str, str] | None:
    stripped = strip_quality_prefix(line)
    if not stripped:
        return None
    lowered = stripped.lower()
    for label in labels:
        label_lower = label.lower()
        if lowered == label_lower:
            return label, ""
        for sep in ("：", ":"):
            prefix = f"{label}{sep}"
            if stripped.startswith(prefix) or lowered.startswith(prefix.lower()):
                return label, stripped[len(prefix):].strip()
    return None


def find_explicit_quality_text(
    lines: list[str],
    labels: list[str],
    stop_headings: set[str],
    all_quality_labels: list[str],
) -> str:
    for index, line in enumerate(lines):
        match = split_quality_label_line(line, labels)
        if match is None:
            continue
        _, inline_text = match
        if inline_text:
            return inline_text
        collected: list[str] = []
        for follow in lines[index + 1:]:
            stripped = follow.strip()
            if not stripped:
                continue
            if stripped in stop_headings:
                break
            if split_quality_label_line(stripped, all_quality_labels) is not None:
                break
            collected.append(stripped)
        return "\n".join(collected).strip()
    return ""


def has_any_explicit_quality_label(lines: list[str], language: str) -> bool:
    labels = [label for group in quality_labels(language).values() for label in group]
    return any(split_quality_label_line(line, labels) is not None for line in lines)


def legacy_quality_search_space(lines: list[str], language: str) -> str:
    if language == "zh-CN":
        headings = ZH_HEADINGS | {"关键代码上下文", "验证环境关键文件"}
        analysis_heading = "漏洞分析"
        risk_heading = "漏洞危险性评估"
    else:
        headings = EN_HEADINGS | {"Key Code Context", "Verification Environment Key Files"}
        analysis_heading = "Vulnerability Analysis"
        risk_heading = "Risk Assessment"
    analysis = section_text(lines, analysis_heading, headings)
    risk = section_text(lines, risk_heading, headings)
    return "\n".join(part for part in (risk, analysis) if part).strip() or "\n".join(lines)


def find_legacy_quality_text(lines: list[str], language: str, kind: str) -> str:
    search_space = legacy_quality_search_space(lines, language)
    markers = legacy_quality_markers(language)[kind]
    hits = []
    lowered_markers = [marker.lower() for marker in markers]
    for raw in search_space.splitlines():
        line = raw.strip()
        if not line:
            continue
        lowered = line.lower()
        if any(marker in line or marker in lowered for marker in markers + lowered_markers):
            hits.append(line)
    if kind == "attacker" and not hits:
        for raw in search_space.splitlines():
            line = raw.strip()
            if language == "zh-CN" and "攻击者" in line and ("可控" in line or "认证" in line or "权限" in line):
                hits.append(line)
            elif language == "en-US" and "attacker" in line.lower() and (
                "control" in line.lower() or "auth" in line.lower() or "privilege" in line.lower()
            ):
                hits.append(line)
    if kind == "server" and not hits:
        for raw in search_space.splitlines():
            line = raw.strip()
            lowered = line.lower()
            if language == "zh-CN" and ("默认" in line or "配置" in line or "环境" in line):
                hits.append(line)
            elif language == "en-US" and ("default" in lowered or "configuration" in lowered or "environment" in lowered):
                hits.append(line)
    if kind == "impact" and not hits:
        patterns = ZH_IMPACT_PATTERNS if language == "zh-CN" else EN_IMPACT_PATTERNS
        for raw in search_space.splitlines():
            line = raw.strip()
            if any(pattern.search(line) for pattern in patterns):
                hits.append(line)
    return "\n".join(hits[:3]).strip()


def placeholder_only_quality_text(text: str, language: str) -> bool:
    normalized = normalize_token(text)
    if not normalized:
        return True
    placeholders = ZH_QUALITY_PLACEHOLDERS if language == "zh-CN" else EN_QUALITY_PLACEHOLDERS
    return normalized in {normalize_token(item) for item in placeholders}


def validate_quality_text_meaningful(text: str, language: str, kind: str) -> None:
    label = quality_display_name(language, kind)
    if not text.strip():
        fail(f"report quality gate missing {label}; add a non-empty {label} section")
    if placeholder_only_quality_text(text, language):
        fail(f"report quality gate {label} is placeholder-only; replace it with concrete reviewer-facing conditions")


def validate_quality_impact_wording(text: str, language: str) -> None:
    patterns = ZH_IMPACT_PATTERNS if language == "zh-CN" else EN_IMPACT_PATTERNS
    if not any(pattern.search(text) for pattern in patterns):
        label = quality_display_name(language, "impact")
        fail(
            f"report quality gate {label} must mention a concrete CIA impact or equivalent "
            "such as data exposure, unauthorized modification, DoS, SSRF/internal access, or information disclosure"
        )


def validate_quality_gate(lines: list[str], language: str) -> None:
    labels_by_kind = quality_labels(language)
    all_quality_labels = [label for group in labels_by_kind.values() for label in group]
    stop_headings = (
        ZH_HEADINGS | {"关键代码上下文", "验证环境关键文件"}
        if language == "zh-CN"
        else EN_HEADINGS | {"Key Code Context", "Verification Environment Key Files"}
    )
    strict_explicit_mode = has_any_explicit_quality_label(lines, language)

    extracted: dict[str, str] = {}
    for kind, labels in labels_by_kind.items():
        explicit_text = find_explicit_quality_text(lines, labels, stop_headings, all_quality_labels)
        extracted[kind] = explicit_text if strict_explicit_mode else (explicit_text or find_legacy_quality_text(lines, language, kind))
        validate_quality_text_meaningful(extracted[kind], language, kind)

    validate_quality_impact_wording(extracted["impact"], language)


def real_world_heading_aliases(language: str) -> set[str]:
    return ZH_REAL_WORLD_HEADINGS if language == "zh-CN" else EN_REAL_WORLD_HEADINGS


def line_contains_heading_alias(line: str, aliases: set[str]) -> bool:
    stripped = re.sub(r"^[#>\-\s\d.、)）]+", "", str(line).strip())
    lowered = stripped.lower()
    return any(alias.lower() in lowered for alias in aliases)


def exploitability_stop_headings(language: str) -> set[str]:
    common = {
        "关键代码上下文",
        "验证环境关键文件",
        "补充材料说明",
        "结论",
        "Attachment Usage",
        "Bundled Materials",
        "Conclusion",
        "Key Code Context",
        "Key Verification Environment Files",
    }
    return (ZH_HEADINGS | common) if language == "zh-CN" else (EN_HEADINGS | common)


def extract_real_world_sections(lines: list[str], language: str) -> list[str]:
    aliases = real_world_heading_aliases(language)
    stops = exploitability_stop_headings(language)
    sections: list[str] = []
    for index, line in enumerate(lines):
        if not line_contains_heading_alias(line, aliases):
            continue
        collected: list[str] = []
        for follow in lines[index + 1:]:
            stripped = follow.strip()
            if not stripped:
                continue
            markdown_heading = re.match(r"^#{1,4}\s+", stripped)
            if stripped in stops or (markdown_heading and not line_contains_heading_alias(stripped, aliases)):
                break
            collected.append(stripped)
        sections.append("\n".join(collected).strip())
    return sections


def exploitability_markers(language: str) -> dict[str, tuple[str, ...]]:
    return ZH_EXPLOITABILITY_MARKERS if language == "zh-CN" else EN_EXPLOITABILITY_MARKERS


def exploitability_section_placeholder_only(text: str, language: str) -> bool:
    normalized = normalize_token(text)
    if not normalized or len(normalized) < 24:
        return True
    placeholders = ZH_QUALITY_PLACEHOLDERS if language == "zh-CN" else EN_QUALITY_PLACEHOLDERS
    placeholder_tokens = {normalize_token(item) for item in placeholders}
    return normalized in placeholder_tokens


def exploitability_missing_kinds(text: str, language: str) -> list[str]:
    lowered = text.lower()
    missing: list[str] = []
    for kind, markers in exploitability_markers(language).items():
        if not any(marker.lower() in lowered for marker in markers):
            missing.append(kind)
    return missing


def is_concrete_exploitability_text(text: str, language: str) -> bool:
    if exploitability_section_placeholder_only(text, language):
        return False
    return not exploitability_missing_kinds(text, language)


def legacy_exploitability_search_text(lines: list[str], supplement_text: str, language: str) -> str:
    if language == "zh-CN":
        headings = ZH_HEADINGS | {"关键代码上下文", "验证环境关键文件"}
        analysis = section_text(lines, "漏洞分析", headings)
        reproduction = section_text(lines, "漏洞复现", headings)
    else:
        headings = EN_HEADINGS | {"Key Code Context", "Key Verification Environment Files"}
        analysis = section_text(lines, "Vulnerability Analysis", headings)
        reproduction = section_text(lines, "Reproduction", headings)
    return "\n".join(part for part in (analysis, reproduction, supplement_text) if part).strip()


def validate_real_world_exploitability(lines: list[str], supplement_text: str, language: str) -> str:
    sections = extract_real_world_sections(lines, language)
    sections.extend(extract_real_world_sections(supplement_text.splitlines(), language))
    for section in sections:
        if is_concrete_exploitability_text(section, language):
            return section

    if sections:
        display = EXPLOITABILITY_DISPLAY[language]
        missing = exploitability_missing_kinds("\n".join(sections), language)
        missing_text = ", ".join(display[kind] for kind in missing) if missing else "concrete non-placeholder text"
        fail(
            "real-world exploitability section is too thin or placeholder-only; "
            f"add reviewer-facing detail for: {missing_text}"
        )

    legacy_text = legacy_exploitability_search_text(lines, supplement_text, language)
    if is_concrete_exploitability_text(legacy_text, language):
        warn(
            "real-world exploitability explanation was accepted from existing analysis/reproduction wording; "
            "prefer a dedicated Real-World Exploitability section in future bundles"
        )
        return legacy_text

    display = EXPLOITABILITY_DISPLAY[language]
    missing = exploitability_missing_kinds(legacy_text, language)
    missing_text = ", ".join(display[kind] for kind in missing) if missing else "concrete non-placeholder text"
    fail(
        "missing real-world exploitability explanation; add a concise section explaining "
        f"{missing_text}"
    )


def validate_poc_attacker_capability_boundary(material_text: str, exploitability_text: str, language: str) -> None:
    matches = sorted({match.group(0) for pattern in STRONG_ATTACKER_CONTROL_PATTERNS for match in pattern.finditer(material_text)})
    if not matches:
        return
    lowered = exploitability_text.lower()
    boundary_markers = (
        ("边界", "限定", "不声称", "未验证", "不扩大", "更强")
        if language == "zh-CN"
        else ("boundary", "limited", "not claim", "not verified", "stronger", "unrelated capability")
    )
    if not any(marker.lower() in lowered for marker in boundary_markers):
        fail(
            "strong-attacker-control PoC wording appears without an explicit target-component boundary explanation: "
            + ", ".join(matches[:5])
        )


def severity_cn_from_any(value: object) -> str:
    text = str(value or "").strip()
    lowered = text.lower()
    if "critical" in lowered or "严重" in text:
        return "严重"
    if "high" in lowered or "高危" in text:
        return "高危"
    if "medium" in lowered or "中危" in text:
        return "中危"
    if "low" in lowered or "低危" in text:
        return "低危"
    return "中危"


def severity_en_from_any(value: object) -> str:
    text = str(value or "").strip()
    lowered = text.lower()
    if "critical" in lowered or "严重" in text:
        return "Critical"
    if "high" in lowered or "高危" in text:
        return "High"
    if "medium" in lowered or "中危" in text:
        return "Medium"
    if "low" in lowered or "低危" in text:
        return "Low"
    return "Medium"


def severity_label_from_score(score_text: object, language: str) -> str:
    try:
        score = float(str(score_text).strip())
    except ValueError:
        return "Medium" if language == "en-US" else "中危"
    if score >= 9:
        return "Critical" if language == "en-US" else "严重"
    if score >= 7:
        return "High" if language == "en-US" else "高危"
    if score >= 4:
        return "Medium" if language == "en-US" else "中危"
    return "Low" if language == "en-US" else "低危"


def resolve_findings_path(bundle_dir: Path) -> Path | None:
    local_findings = bundle_dir / "findings.json"
    if local_findings.exists():
        return local_findings
    confirmed_findings = bundle_dir.parent / "findings.json"
    if confirmed_findings.exists():
        return confirmed_findings
    return None


def load_selected_finding(findings_path: Path, bundle_name: str, workspace_dir: Path) -> tuple[dict[str, object], dict[str, object]]:
    try:
        data = json.loads(findings_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        fail(f"bundle findings.json is unreadable: {findings_path} ({exc})")

    defaults: dict[str, object] = {}
    findings: list[dict[str, object]] = []
    if isinstance(data, dict) and isinstance(data.get("findings"), list):
        defaults = {k: v for k, v in data.items() if k != "findings"}
        findings = [item for item in data["findings"] if isinstance(item, dict)]
    elif isinstance(data, dict):
        findings = [data]
    elif isinstance(data, list):
        findings = [item for item in data if isinstance(item, dict)]
    else:
        fail(f"bundle findings.json must be a JSON object or array: {findings_path}")

    if not findings:
        fail(f"bundle findings.json does not contain any finding objects: {findings_path}")
    if findings_path.parent.name == bundle_name and len(findings) != 1:
        fail(
            "per-bundle findings.json must describe exactly one confirmed vulnerability. "
            f"Found {len(findings)} findings; split them into one bundle per finding."
        )

    selected = None
    for item in findings:
        slug = str(item.get("slug", "")).strip()
        report_file = Path(str(item.get("report_file", "")).strip() or "x").stem
        filename = Path(str(item.get("filename", "")).strip() or "x").stem
        if slug == bundle_name or report_file == bundle_name or filename == bundle_name:
            selected = item
            break
    if selected is None and len(findings) == 1:
        selected = findings[0]
    if selected is None:
        fail(f"could not match {findings_path.name} to bundle directory {bundle_name}")

    return defaults, selected


def load_bundle_identity(findings_path: Path, bundle_name: str, workspace_dir: Path) -> tuple[str, list[str], dict[str, object]]:
    defaults, selected = load_selected_finding(findings_path, bundle_name, workspace_dir)

    project_name = (
        str(selected.get("project_name") or selected.get("project") or defaults.get("project_name") or defaults.get("project")).strip()
    )
    if not project_name:
        project_name = workspace_dir.parent.name
    project_name = project_name.split("/")[-1]

    raw_title_tokens = [
        str(selected.get("title", "")).strip(),
        str(selected.get("title_zh", "")).strip(),
        str(selected.get("title_en", "")).strip(),
        str(selected.get("vulnerability_name", "")).strip(),
        str(selected.get("vulnerability_name_zh", "")).strip(),
        str(selected.get("vulnerability_name_en", "")).strip(),
        str(selected.get("vuln_type", "")).strip(),
        str(selected.get("vuln_type_zh", "")).strip(),
        str(selected.get("vuln_type_en", "")).strip(),
        str(selected.get("type", "")).strip(),
        str(selected.get("vulnerability_type", "")).strip(),
    ]
    title_tokens = [token for token in raw_title_tokens if token]
    return project_name, title_tokens, selected


def looks_like_short_vulnerability_name(value: object) -> bool:
    text = str(value or "").strip()
    if not text or "\n" in text or "\r" in text:
        return False
    lowered = text.lower()
    if "漏洞报告" in text or "vulnerability report" in lowered:
        return False
    if len(text) > 80:
        return False
    if not any("\u4e00" <= ch <= "\u9fff" for ch in text) and len(text.split()) > 12:
        return False
    return True


def validate_finding_vulnerability_name(finding: dict[str, object]) -> None:
    explicit_names = [
        finding.get("vulnerability_name"),
        finding.get("vulnerability_name_zh"),
        finding.get("vulnerability_name_en"),
    ]
    if any(looks_like_short_vulnerability_name(value) for value in explicit_names):
        return
    explicit_types = [
        finding.get("vuln_type"),
        finding.get("vuln_type_zh"),
        finding.get("vuln_type_en"),
        finding.get("vulnerability_type"),
        finding.get("type"),
    ]
    if any(looks_like_short_vulnerability_name(value) for value in explicit_types):
        return
    title = str(finding.get("title") or finding.get("title_zh") or finding.get("title_en") or "").strip()
    fail(
        "findings.json must include vulnerability_name/vulnerability_name_zh/vulnerability_name_en "
        "or a short vuln_type/vuln_type_zh/vuln_type_en; do not derive bundle identity from a long title"
        + (f": {title[:140]}" if title else "")
    )


def nested_value(data: object, dotted_key: str) -> object:
    current = data
    for part in dotted_key.split("."):
        if not isinstance(current, dict):
            return None
        current = current.get(part)
    return current


def boolish_false(value: object) -> bool:
    if value is False:
        return True
    text = str(value or "").strip().lower()
    return text in {"false", "no", "0", "mismatch", "not_matched", "runtime_only"}


def detect_runtime_mismatch(defaults: dict[str, object], finding: dict[str, object], workspace_dir: Path) -> bool:
    containers: list[object] = [
        defaults,
        finding,
        finding.get("runtime_scope") if isinstance(finding.get("runtime_scope"), dict) else {},
        finding.get("verification_evidence") if isinstance(finding.get("verification_evidence"), dict) else {},
        defaults.get("metadata") if isinstance(defaults.get("metadata"), dict) else {},
    ]
    for data in containers:
        if not isinstance(data, dict):
            continue
        for key in ("source_runtime_match", "runtime_source_match"):
            if key in data and boolish_false(data.get(key)):
                return True
        for key in ("verified_runtime_only", "runtime_scope"):
            value = data.get(key)
            if value is True or (isinstance(value, str) and value.strip()):
                return True

    for rel in ("workspace-report.md", "audit-report.md", "audit-research-report.md", "attack-surface.md"):
        path = workspace_dir / rel
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        if any(pattern.search(text) for pattern in RUNTIME_MISMATCH_TEXT_PATTERNS):
            return True
    return False


def runtime_scope_statement(defaults: dict[str, object], finding: dict[str, object], evidence: dict[str, object]) -> str:
    values = [
        finding.get("runtime_scope"),
        finding.get("verified_runtime_only"),
        finding.get("verified_runtime_version"),
        finding.get("runtime_version"),
        finding.get("runtime_image_digest"),
        finding.get("runtime_commit"),
        nested_value(defaults, "metadata.runtime_scope"),
        nested_value(defaults, "metadata.verified_runtime_only"),
        nested_value(defaults, "metadata.runtime_version"),
        nested_value(defaults, "metadata.docker_image"),
        evidence.get("runtime_scope"),
        evidence.get("verified_runtime_only"),
        evidence.get("docker_image"),
    ]
    return " ".join(str(value).strip() for value in values if str(value or "").strip())


def affected_version_text(defaults: dict[str, object], finding: dict[str, object]) -> str:
    impact = finding.get("impact") if isinstance(finding.get("impact"), dict) else {}
    values = [
        finding.get("affected_versions"),
        finding.get("affected_versions_zh"),
        finding.get("affected_versions_en"),
        impact.get("affected_versions") if isinstance(impact, dict) else "",
        impact.get("affected_versions_zh") if isinstance(impact, dict) else "",
        impact.get("affected_versions_en") if isinstance(impact, dict) else "",
        nested_value(defaults, "metadata.affected_versions"),
    ]
    return "\n".join(str(value).strip() for value in values if str(value or "").strip())


def validate_runtime_scope(
    defaults: dict[str, object],
    finding: dict[str, object],
    evidence: dict[str, object],
    workspace_dir: Path,
) -> None:
    if not detect_runtime_mismatch(defaults, finding, workspace_dir):
        return
    versions = affected_version_text(defaults, finding)
    for pattern in BROAD_AFFECTED_VERSION_PATTERNS:
        if pattern.search(versions):
            fail(
                "source/runtime mismatch detected; affected_versions must not overclaim broad unsupported scope "
                f"({pattern.pattern}). Limit the confirmed bundle to the verified runtime version/image/commit."
            )
    scope = runtime_scope_statement(defaults, finding, evidence)
    if not scope:
        fail(
            "source/runtime mismatch detected; include verified_runtime_only/runtime_scope or the verified "
            "runtime version, image digest, or commit before validating the confirmed bundle"
        )


def validate_severity_consistency(
    bundle_dir: Path,
    docx_path: Path,
    finding: dict[str, object],
    language: str,
) -> None:
    cvss = finding.get("cvss")
    if isinstance(cvss, dict):
        score_text = str(cvss.get("score") or "").strip()
        severity_text = str(cvss.get("severity") or "").strip()
    else:
        score_text = str(finding.get("cvss4_score") or finding.get("cvss_score") or "").strip()
        severity_text = str(finding.get("severity_cn") if language == "zh-CN" else finding.get("severity") or "").strip()
    if not score_text and not severity_text:
        return

    expected = severity_label_from_score(score_text, language) if score_text else (
        severity_en_from_any(severity_text) if language == "en-US" else severity_cn_from_any(severity_text)
    )
    expected_norm = normalize_token(expected)
    if not expected_norm:
        return

    if severity_text:
        cvss_declared = severity_en_from_any(severity_text) if language == "en-US" else severity_cn_from_any(severity_text)
        if normalize_token(cvss_declared) != expected_norm:
            fail(
                "cvss.severity does not match the severity implied by cvss.score: "
                f"score={score_text}, severity={severity_text}"
            )

    explicit_severity = str(finding.get("severity_cn") if language == "zh-CN" else finding.get("severity") or "").strip()
    if explicit_severity:
        explicit_label = severity_en_from_any(explicit_severity) if language == "en-US" else severity_cn_from_any(explicit_severity)
        if normalize_token(explicit_label) != expected_norm:
            fail(
                "findings severity label does not match the severity implied by CVSS: "
                f"declared={explicit_severity}, expected={expected}"
            )

    if expected_norm not in normalize_token(docx_path.stem):
        fail(
            "report filename severity does not match the severity implied by CVSS: "
            f"expected {expected}, got {docx_path.stem}"
        )
    if expected_norm not in normalize_token(bundle_dir.name):
        fail(
            "bundle directory severity does not match the severity implied by CVSS: "
            f"expected {expected}, got {bundle_dir.name}"
        )

    explicit_filename = str(finding.get("filename") or finding.get("report_file") or "").strip()
    if explicit_filename and expected_norm not in normalize_token(Path(explicit_filename).stem):
        fail(
            "findings filename does not match the severity implied by CVSS: "
            f"expected {expected}, got {explicit_filename}"
        )


def cvss_vector_from_finding(finding: dict[str, object]) -> str:
    cvss = finding.get("cvss")
    if isinstance(cvss, dict):
        return str(cvss.get("vector") or "").strip()
    return str(finding.get("cvss_vector") or finding.get("cvss31_vector") or finding.get("cvss4_vector") or "").strip()


def cvss_requires_privileges(vector: str) -> bool:
    return bool(re.search(r"(?:^|/)PR:[LH](?:/|$)", str(vector or ""), re.IGNORECASE))


def unauthenticated_claim_matches(text: str) -> list[str]:
    return [match.group(0) for pattern in UNAUTHENTICATED_TITLE_PATTERNS for match in pattern.finditer(text or "")]


def contains_authenticated_flow_markers(text: str) -> bool:
    return any(pattern.search(text or "") for pattern in AUTHENTICATED_FLOW_MARKERS)


def finding_reproduction_text(finding: dict[str, object]) -> str:
    parts: list[str] = []
    raw_steps = finding.get("reproduction")
    if isinstance(raw_steps, list):
        for step in raw_steps:
            if isinstance(step, dict):
                parts.extend(str(value) for value in step.values())
            else:
                parts.append(str(step))
    for key in ("reproduction", "steps", "proof", "poc", "verification_evidence"):
        value = finding.get(key)
        if isinstance(value, (dict, list)):
            parts.append(json.dumps(value, ensure_ascii=False, sort_keys=True))
        elif value is not None:
            parts.append(str(value))
    return "\n".join(parts)


def validate_title_auth_cvss_consistency(
    bundle_dir: Path,
    docx_path: Path,
    docx_lines: list[str],
    finding: dict[str, object],
) -> None:
    title_values = [
        docx_lines[0] if docx_lines else "",
        bundle_dir.name,
        docx_path.stem,
        str(finding.get("title") or ""),
        str(finding.get("title_zh") or ""),
        str(finding.get("title_en") or ""),
        str(finding.get("filename") or ""),
        str(finding.get("report_file") or ""),
    ]
    title_text = "\n".join(value for value in title_values if str(value).strip())
    claims = unauthenticated_claim_matches(title_text)
    if not claims:
        return

    vector = cvss_vector_from_finding(finding)
    auth_markers = contains_authenticated_flow_markers(finding_reproduction_text(finding) + "\n" + "\n".join(docx_lines))
    if cvss_requires_privileges(vector) or auth_markers:
        reason = f"CVSS vector {vector or '<missing>'} requires privileges" if cvss_requires_privileges(vector) else "reproduction uses authenticated-flow markers"
        fail(
            "title/CVSS/auth consistency failure: report title, DOCX first line, bundle directory, "
            "or findings title must not claim unauthenticated/pre-auth reachability when "
            f"{reason}. Offending title term(s): {', '.join(sorted(set(claims)))}"
        )


def validate_bundle_identity(bundle_dir: Path, docx_lines: list[str], workspace_dir: Path) -> None:
    findings_path = resolve_findings_path(bundle_dir)
    if findings_path is None:
        return
    project_name, title_tokens, _ = load_bundle_identity(findings_path, bundle_dir.name, workspace_dir)
    combined_text = "\n".join(docx_lines)
    first_line = docx_lines[0] if docx_lines else ""
    normalized_first_line = normalize_token(first_line)
    normalized_combined = normalize_token(combined_text)
    normalized_project = normalize_token(project_name)

    if normalized_project and normalized_project not in normalized_combined:
        fail(f"report content does not mention the expected project identity from findings.json: {project_name}")
    if normalized_project and normalized_project not in normalized_first_line:
        fail(f"report title does not match the expected project identity from findings.json: {project_name}")

    normalized_titles = [normalize_token(token) for token in title_tokens if normalize_token(token)]
    if normalized_titles and not any(token in normalized_combined for token in normalized_titles):
        fail("report content does not match the expected vulnerability identity from findings.json")
    if normalized_titles and not any(token in normalized_first_line for token in normalized_titles):
        fail("report title does not match the expected vulnerability identity from findings.json")

    generic_title_tokens = {normalize_token("安全漏洞"), normalize_token("Vulnerability")}
    if normalized_titles and normalize_token(first_line) in generic_title_tokens:
        fail("report title is generic; use the finding-specific vulnerability identity from findings.json")


def validate_material_identity(text: str, label: str, project_name: str, title_tokens: list[str]) -> None:
    normalized_text = normalize_token(text)
    normalized_project = normalize_token(project_name)
    if normalized_project and normalized_project not in normalized_text:
        fail(f"{label} does not mention the expected project identity from findings.json: {project_name}")

    normalized_titles = [normalize_token(token) for token in title_tokens if normalize_token(token)]
    if normalized_titles and not any(token in normalized_text for token in normalized_titles):
        fail(f"{label} does not match the expected vulnerability identity from findings.json")


def validate_language_consistency(
    text: str,
    label: str,
    language: str,
    *,
    expected_markers: list[str],
    opposite_markers: list[str],
) -> None:
    if not any(marker in text for marker in expected_markers):
        fail(f"{label} is missing obvious reviewer-facing markers for {language}")
    if any(marker in text for marker in opposite_markers):
        fail(f"{label} contains reviewer-facing markers from the wrong output language")


def validate_no_absolute_paths(text: str) -> None:
    for pattern in ABSOLUTE_PATH_PATTERNS:
        match = pattern.search(text)
        if match:
            fail(f"found absolute/operator-local path reference: {match.group(0)}")


def validate_no_placeholder_text(text: str, language: str, label: str) -> None:
    for placeholder in PLACEHOLDER_REPORT_TEXT.get(language, []):
        if placeholder in text:
            fail(f"{label} contains placeholder text that must be resolved before validation: {placeholder}")


def collect_json_strings(value: object) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        strings: list[str] = []
        for nested in value.values():
            strings.extend(collect_json_strings(nested))
        return strings
    if isinstance(value, list):
        strings = []
        for nested in value:
            strings.extend(collect_json_strings(nested))
        return strings
    return []


def addendum_placeholder_only(text: str) -> bool:
    body = re.sub(r"```.*?```", "", text, flags=re.DOTALL)
    body = re.sub(r"^[#>\-\s\d.、)）]+", "", body, flags=re.MULTILINE).strip()
    normalized = normalize_token(body)
    if len(normalized) < 80:
        return True
    meaningful_lines = [
        line.strip()
        for line in body.splitlines()
        if line.strip() and not line.strip().startswith("#")
    ]
    if not meaningful_lines:
        return True
    placeholder_lines = sum(
        1 for line in meaningful_lines
        if any(pattern.search(line) for pattern in REVIEWER_PLACEHOLDER_PATTERNS)
    )
    return placeholder_lines == len(meaningful_lines)


def validate_reviewer_evidence_addendum(bundle_dir: Path, language: str) -> str:
    path = bundle_dir / REVIEWER_EVIDENCE_ADDENDUM
    if not path.exists():
        return ""
    text = path.read_text(encoding="utf-8")
    validate_no_absolute_paths(text)
    validate_relative_attachment_refs(text, bundle_dir)
    validate_no_placeholder_text(text, language, "reviewer evidence addendum")
    if addendum_placeholder_only(text):
        fail("reviewer-evidence-and-impact.md is placeholder-only; add attacker boundary, evidence, oracles, and replay detail")

    lowered = text.lower()
    missing = [
        group
        for group, markers in REVIEWER_ADDENDUM_MARKERS.items()
        if not any(marker.lower() in lowered for marker in markers)
    ]
    if missing:
        fail(
            "reviewer-evidence-and-impact.md is missing useful reviewer evidence wording for: "
            + ", ".join(missing)
        )
    return text


def reviewer_index_artifact_key(key: str) -> bool:
    normalized = key.strip().lower()
    if normalized in REVIEWER_INDEX_ARTIFACT_KEYS:
        return True
    return normalized.endswith(("_artifact", "_artifacts", "_file", "_files", "_path", "_paths"))


def collect_reviewer_index_artifact_paths(value: object, key: str = "") -> list[tuple[str, str]]:
    paths: list[tuple[str, str]] = []
    if isinstance(value, dict):
        for child_key, child_value in value.items():
            if reviewer_index_artifact_key(str(child_key)):
                for item in collect_json_strings(child_value):
                    if item.strip():
                        paths.append((str(child_key), item.strip()))
            else:
                paths.extend(collect_reviewer_index_artifact_paths(child_value, str(child_key)))
    elif isinstance(value, list):
        for item in value:
            paths.extend(collect_reviewer_index_artifact_paths(item, key))
    return paths


def collect_reviewer_index_oracles(value: object, key: str = "") -> list[str]:
    tokens: list[str] = []
    if isinstance(value, dict):
        for child_key, child_value in value.items():
            normalized = str(child_key).strip().lower()
            if normalized in REVIEWER_INDEX_ORACLE_KEYS:
                tokens.extend(item.strip() for item in collect_json_strings(child_value) if item.strip())
            else:
                tokens.extend(collect_reviewer_index_oracles(child_value, normalized))
    elif isinstance(value, list):
        for item in value:
            tokens.extend(collect_reviewer_index_oracles(item, key))
    return tokens


def collect_reviewer_command(data: dict[str, object]) -> str:
    for key in REVIEWER_COMMAND_KEYS:
        value = str(data.get(key) or "").strip()
        if value:
            return value
    return ""


def validate_reviewer_index_artifact_path(bundle_dir: Path, raw: str, label: str) -> None:
    if not raw:
        fail(f"reviewer evidence index {label} must not be empty")
    if "oss-vulnerability-research" in raw:
        fail(f"reviewer evidence index {label} contains submitter repository path text: {raw}")
    if re.search(r"/tmp/[A-Za-z0-9_.-]+-audit\b", raw):
        fail(f"reviewer evidence index {label} contains submitter-local audit path: {raw}")
    for pattern in ABSOLUTE_PATH_PATTERNS:
        if pattern.search(raw):
            fail(f"reviewer evidence index {label} must be bundle-relative, got operator-local path: {raw}")
    if "://" in raw:
        fail(f"reviewer evidence index {label} must reference a bundle-local file, not a URL: {raw}")
    if re.match(r"^[A-Za-z]:[\\/]", raw):
        fail(f"reviewer evidence index {label} must not use a Windows absolute path: {raw}")
    path = Path(raw)
    if path.is_absolute() or raw.startswith("~"):
        fail(f"reviewer evidence index {label} must be bundle-relative, got absolute path: {raw}")
    if ".." in path.parts:
        fail(f"reviewer evidence index {label} must stay inside the bundle and not use '..': {raw}")
    resolved = (bundle_dir / path).resolve()
    try:
        resolved.relative_to(bundle_dir.resolve())
    except ValueError:
        fail(f"reviewer evidence index {label} escapes the bundle root: {raw}")
    if not resolved.exists():
        fail(f"reviewer evidence index referenced artifact path does not exist: {raw}")


def validate_reviewer_command(bundle_dir: Path, command: str) -> None:
    if not command:
        fail("attachments/reviewer-evidence-index.json must include a bundle-root reviewer replay command")
    if "oss-vulnerability-research" in command or re.search(r"/tmp/[A-Za-z0-9_.-]+-audit\b", command):
        fail("reviewer evidence index replay command must not require the submitter repository or audit workspace")
    for pattern in ABSOLUTE_PATH_PATTERNS:
        if pattern.search(command):
            fail(f"reviewer evidence index replay command must be bundle-root local, got: {command}")
    if re.search(r"(^|[;&|]\s*)cd\s+\.\.", command) or "../" in command:
        fail(f"reviewer evidence index replay command must not climb to a parent repository: {command}")
    script_refs = re.findall(r"(?:^|\s)(?:bash\s+|sh\s+|zsh\s+)?(\./[A-Za-z0-9_.\-/]+\.sh)\b", command)
    if not script_refs:
        fail("reviewer evidence index replay command must call a bundle-root ./run-*.sh helper")
    for ref in script_refs:
        script = (bundle_dir / ref[2:]).resolve()
        try:
            script.relative_to(bundle_dir.resolve())
        except ValueError:
            fail(f"reviewer evidence index replay command script escapes bundle: {ref}")
        if not script.exists():
            fail(f"reviewer evidence index replay command references a missing script: {ref}")


def read_text_source(path: Path) -> str:
    try:
        if path.stat().st_size > 1_500_000:
            return ""
        return path.read_text(encoding="utf-8", errors="ignore")
    except OSError:
        return ""


def reviewer_oracle_source_corpus(
    bundle_dir: Path,
    excluded: set[Path],
    docx_text_value: str,
    note_text: str,
    supplement_text: str,
    addendum_text: str,
    evidence: dict[str, object],
) -> str:
    parts = [
        docx_text_value,
        note_text,
        supplement_text,
        addendum_text,
        json.dumps(evidence, ensure_ascii=False, sort_keys=True),
    ]
    for path in sorted(bundle_dir.rglob("*")):
        if not path.is_file() or path.resolve() in excluded:
            continue
        if path.suffix.lower() not in TEXT_EVIDENCE_SUFFIXES:
            continue
        parts.append(read_text_source(path))
    return "\n".join(part for part in parts if part)


def validate_reviewer_evidence_index(
    bundle_dir: Path,
    evidence: dict[str, object],
    docx_text_value: str,
    note_text: str,
    supplement_text: str,
    addendum_text: str,
) -> tuple[str, dict[str, object] | None]:
    path = bundle_dir / REVIEWER_EVIDENCE_INDEX
    if not path.exists():
        return "", None
    raw_text = path.read_text(encoding="utf-8")
    validate_no_absolute_paths(raw_text)
    if "oss-vulnerability-research" in raw_text:
        fail("attachments/reviewer-evidence-index.json must not contain submitter repository path text")
    if re.search(r"/tmp/[A-Za-z0-9_.-]+-audit\b", raw_text):
        fail("attachments/reviewer-evidence-index.json must not contain submitter-local audit paths")
    for pattern in SUBMITTER_LOCAL_IMAGE_PATTERNS:
        match = pattern.search(raw_text)
        if match:
            fail(f"attachments/reviewer-evidence-index.json references a submitter-local image name: {match.group(0)}")
    try:
        data = json.loads(raw_text)
    except json.JSONDecodeError as exc:
        fail(f"attachments/reviewer-evidence-index.json is invalid JSON: {exc}")
    if not isinstance(data, dict):
        fail("attachments/reviewer-evidence-index.json must contain a JSON object")

    artifact_paths = collect_reviewer_index_artifact_paths(data)
    if not artifact_paths:
        fail("attachments/reviewer-evidence-index.json must map evidence artifacts to bundle-local paths")
    for label, raw in artifact_paths:
        validate_reviewer_index_artifact_path(bundle_dir, raw, label)

    command = collect_reviewer_command(data)
    validate_reviewer_command(bundle_dir, command)

    oracle_tokens = sorted(set(collect_reviewer_index_oracles(data)))
    if str(evidence.get("oracle_token") or "").strip() and not oracle_tokens:
        fail("attachments/reviewer-evidence-index.json must list success oracle tokens")
    corpus = reviewer_oracle_source_corpus(
        bundle_dir,
        {path.resolve()},
        docx_text_value,
        note_text,
        supplement_text,
        addendum_text,
        evidence,
    )
    corpus_lower = corpus.lower()
    for token in oracle_tokens:
        if token not in corpus and token.lower() not in corpus_lower:
            suffix = re.sub(r"^[A-Za-z][A-Za-z _-]{1,40}:\s*", "", token).strip()
            if suffix and len(suffix) >= 8 and (suffix in corpus or suffix.lower() in corpus_lower):
                warn(
                    "reviewer evidence index oracle token was accepted by matching its concrete value "
                    f"without the display label: {token}"
                )
                continue
            fail(
                "reviewer evidence index success oracle token is not present in scripts, evidence, "
                f"supplements, addendum, or verification-evidence.json: {token}"
            )

    for label, raw in artifact_paths:
        if "video" not in label.lower() and Path(raw).suffix.lower() not in VIDEO_SUFFIXES:
            continue
        if Path(raw).suffix.lower() not in VIDEO_SUFFIXES:
            fail(f"reviewer evidence index video reference must point to a media file: {raw}")

    return raw_text, data


def material_has_pattern(patterns: list[re.Pattern[str]], text: str) -> bool:
    return any(pattern.search(text) for pattern in patterns)


def validate_no_submitter_local_images(text: str) -> None:
    for pattern in SUBMITTER_LOCAL_IMAGE_PATTERNS:
        match = pattern.search(text)
        if match:
            fail(
                "final reviewer materials reference a submitter-local Docker image name; "
                f"use a bundle-local fixture or public pinned/base image instead: {match.group(0)}"
            )


def warn_remote_latest_without_fixture_rationale(text: str) -> None:
    if not REMOTE_LATEST_IMAGE_PATTERN.search(text):
        return
    lowered = text.lower()
    fallback_markers = (
        "bundle-local",
        "attachments/fixture",
        "uses_bundle_local_fixture",
        "requires_librespeed_latest\": false",
        "requires latest: false",
        "不依赖旧",
        "不依赖不可拉取",
        "内置 fixture",
        "本目录内置",
    )
    if any(marker.lower() in lowered for marker in fallback_markers):
        return
    warn(
        "reviewer replay appears to reference a remote :latest image without an obvious bundle-local fallback "
        "or fixture rationale"
    )


def has_fixture_provenance(text: str) -> bool:
    lowered = text.lower()
    context_blocks = [
        block
        for block in split_markdown_paragraphs(text) + text.splitlines()
        if material_has_pattern(FIXTURE_HINT_PATTERNS, block)
    ]
    context = "\n".join(context_blocks).lower()
    structured_index_provenance = (
        "source_evidence" in lowered
        and ("fixture_files" in lowered or "vendor_fixture" in lowered or "uses_bundle_local_vendor" in lowered)
    )
    source_grounded = any(
        marker in context
        for marker in (
            "source file",
            "vulnerable source",
            "vulnerable pattern",
            "source-grounded",
            "preserve",
            "vendored",
            "attachments/vendor",
            "original source",
            "源码",
            "原始源码",
            "危险源码",
            "漏洞模式",
            "保留",
            "sed",
            "entrypoint",
        )
    ) or structured_index_provenance
    sufficient = any(
        marker in context
        for marker in (
            "sufficient",
            "reproduce the vulnerability boundary",
            "minimal fixture",
            "fixture",
            "vendored",
            "attachments/vendor",
            "足以",
            "复现边界",
            "最小复现",
            "内置文件",
            "只依赖",
        )
    ) or structured_index_provenance
    nonclaim = any(marker in lowered for marker in NONCLAIM_MARKERS)
    return source_grounded and sufficient and nonclaim


def missing_library_boundary_kinds(text: str) -> list[str]:
    lowered = text.lower()
    checks = {
        "public API/function boundary": (
            "public api",
            " api",
            "function",
            "method",
            "q.set",
            "q.get",
            "q.post",
            "q.invoke",
            "函数",
            "方法",
            "api",
            "调用",
        ),
        "attacker-controlled argument": (
            "attacker-controlled",
            "attacker controlled",
            "attacker controls",
            "key",
            "name",
            "argument",
            "filename",
            "metadata",
            "source map",
            "config",
            "攻击者可控",
            "字段名",
            "属性名",
            "参数",
            "元数据",
        ),
        "consumer application pattern": (
            "consumer",
            "consuming application",
            "application-level",
            "upper-layer",
            "调用方",
            "消费方",
            "上层应用",
            "业务应用",
            "桥接",
        ),
        "library-local versus application-level boundary": (
            "library-local",
            "application-level",
            "library itself",
            "does not expose",
            "no direct network",
            "object.prototype",
            "global",
            "local",
            "库本身",
            "不提供网络入口",
            "目标对象",
            "全局",
            "局部",
            "不声称",
        ),
    }
    return [
        label
        for label, markers in checks.items()
        if not any(marker in lowered for marker in markers)
    ]


def validate_reviewer_story_gates(text: str) -> None:
    validate_no_submitter_local_images(text)
    warn_remote_latest_without_fixture_rationale(text)
    if material_has_pattern(FIXTURE_HINT_PATTERNS, text) and not has_fixture_provenance(text):
        fail(
            "fixture-based or vendored-source replay needs a source-grounded provenance explanation: "
            "state what source/pattern is preserved, why the fixture is sufficient, and what stronger impact is not claimed"
        )
    if material_has_pattern(LIBRARY_HINT_PATTERNS, text):
        missing = missing_library_boundary_kinds(text)
        if missing:
            fail(
                "library/package report is missing consumer application boundary wording for: "
                + ", ".join(missing)
            )


def has_nonclaim_context(text: str, start: int, end: int) -> bool:
    window = text[max(0, start - 120): min(len(text), end + 120)].lower()
    return any(marker.lower() in window for marker in NONCLAIM_MARKERS)


def claim_matches_without_nonclaim(text: str, patterns: tuple[str, ...]) -> list[str]:
    matches: list[str] = []
    for pattern_text in patterns:
        pattern = re.compile(pattern_text, re.IGNORECASE)
        for match in pattern.finditer(text):
            if not has_nonclaim_context(text, match.start(), match.end()):
                matches.append(match.group(0))
    return matches


def validate_claim_oracle_consistency(text: str) -> None:
    lowered = text.lower()
    webshell_claims = claim_matches_without_nonclaim(
        text,
        (
            r"\bwebshell\b",
            r"\bHTTP\s+(?:command execution|RCE)\b",
            r"HTTP\s*webshell",
            r"webshell\s*命令执行",
            r"HTTP\s*命令执行",
        ),
    )
    if webshell_claims and not any(
        marker.lower() in lowered
        for marker in (
            "webshell_confirmed",
            "code_execution_confirmed",
            "uid=",
            "www-data",
            "whoami",
            "id output",
            "id 输出",
            "命令输出",
        )
    ):
        fail(
            "report claims webshell or HTTP command execution, but reviewer materials do not show a matching oracle"
        )

    dos_claims = claim_matches_without_nonclaim(
        text,
        (
            r"已验证[^。\n]{0,40}(?:拒绝服务|DoS)",
            r"(?:拒绝服务|DoS)[^。\n]{0,40}(?:已确认|攻击成功|成功触发)",
            r"\bDoS_CONFIRMED\b",
            r"\bdenial of service confirmed\b",
            r"\bcauses denial of service\b",
        ),
    )
    if dos_claims and not any(
        marker.lower() in lowered
        for marker in ("dos_confirmed", "crash", "崩溃", "oom", "timeout", "resource exhausted", "拒绝服务成功")
    ):
        fail("report claims confirmed DoS, but reviewer materials do not show a direct DoS oracle")

    broad_claims = claim_matches_without_nonclaim(
        text,
        (
            r"container escape",
            r"容器逃逸",
            r"host\s+RCE",
            r"宿主机[^。\n]{0,30}(?:代码执行|命令执行|RCE)",
            r"unauthenticated public exploitability",
            r"无需认证[^。\n]{0,30}直接触发",
            r"匿名[^。\n]{0,30}直接触发",
        ),
    )
    if broad_claims:
        fail(
            "report claims container escape, host RCE, or unauthenticated public exploitability without an explicit non-claim boundary"
        )


def explicit_body_severity_labels(lines: list[str], language: str) -> list[str]:
    labels: list[str] = []
    patterns = (
        (
            r"(?:等级判定|风险等级|严重性|漏洞危险性|危险等级|最终评级)\s*[:：]\s*(严重|高危|中危|低危|Critical|High|Medium|Low)",
            r"(?:Severity|Risk rating|Final rating|Base severity)\s*[:：]\s*(Critical|High|Medium|Low|严重|高危|中危|低危)",
        )
        if language == "zh-CN"
        else (
            r"(?:Severity|Risk rating|Final rating|Base severity)\s*[:：]\s*(Critical|High|Medium|Low|严重|高危|中危|低危)",
            r"(?:等级判定|风险等级|严重性|最终评级)\s*[:：]\s*(严重|高危|中危|低危|Critical|High|Medium|Low)",
        )
    )
    for line in lines:
        for pattern in patterns:
            match = re.search(pattern, line, re.IGNORECASE)
            if match:
                labels.append(match.group(1))
    return labels


def validate_report_body_severity_consistency(
    bundle_dir: Path,
    docx_path: Path,
    lines: list[str],
    finding: dict[str, object],
    language: str,
) -> None:
    cvss = finding.get("cvss")
    score_text = str(cvss.get("score") or "").strip() if isinstance(cvss, dict) else str(
        finding.get("cvss4_score") or finding.get("cvss_score") or ""
    ).strip()
    if not score_text:
        return
    expected = severity_label_from_score(score_text, language)
    expected_norm = normalize_token(expected)
    if not expected_norm:
        return
    for label in explicit_body_severity_labels(lines, language):
        normalized = severity_en_from_any(label) if language == "en-US" else severity_cn_from_any(label)
        if normalize_token(normalized) != expected_norm:
            fail(
                "report body severity does not match the severity implied by CVSS: "
                f"score={score_text}, expected={expected}, body={label}, bundle={bundle_dir.name}, docx={docx_path.name}"
            )


def validate_bundle_cleanliness(bundle_dir: Path) -> None:
    for child in bundle_dir.iterdir():
        if child.is_dir() and child.name in FORBIDDEN_BUNDLE_DIRS:
            if child.name == "evidence":
                fail("final confirmed bundle must not use evidence/ as a delivery directory; move evidence files under attachments/")
            fail(f"final confirmed bundle must not contain runtime or source-control directory: {child.name}/")


def validate_bundle_relative_file(value: object, bundle_dir: Path, label: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        if label == "poc_path":
            fail(
                "verification-evidence.json poc_path must not be empty; set "
                "findings[].verification_evidence.poc_path to a bundle-relative "
                "PoC file under attachments/."
            )
        fail(f"verification-evidence.json {label} must not be empty")
    if raw.startswith("file://"):
        fail(f"verification-evidence.json {label} must be bundle-relative, got file URI: {raw}")
    if re.match(r"^[A-Za-z]:[\\/]", raw):
        fail(f"verification-evidence.json {label} must be bundle-relative, got Windows absolute path: {raw}")
    path = Path(raw)
    if path.is_absolute():
        fail(f"verification-evidence.json {label} must be bundle-relative, got absolute path: {raw}")
    if ".." in path.parts:
        fail(f"verification-evidence.json {label} must not escape the bundle with '..': {raw}")
    # Path.resolve() follows symlinks; the relative_to() check below therefore
    # rejects symlink escapes as well as ordinary path traversal.
    resolved = (bundle_dir / path).resolve()
    try:
        resolved.relative_to(bundle_dir.resolve())
    except ValueError:
        fail(f"verification-evidence.json {label} escapes the bundle root: {raw}")
    if not resolved.exists():
        fail(f"verification-evidence.json {label} does not exist inside bundle: {raw}")
    if not resolved.is_file():
        fail(f"verification-evidence.json {label} must point to a file: {raw}")
    return path.as_posix()


def validate_verification_evidence(bundle_dir: Path, finding: dict[str, object] | None = None) -> dict[str, object]:
    evidence_path = bundle_dir / "verification-evidence.json"
    if not evidence_path.exists():
        fail("confirmed bundle must include verification-evidence.json")
    try:
        evidence = json.loads(evidence_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        fail(f"verification-evidence.json is unreadable: {exc}")
    if not isinstance(evidence, dict):
        fail("verification-evidence.json must contain a JSON object")

    missing = sorted(field for field in REQUIRED_VERIFICATION_FIELDS if field not in evidence)
    if missing:
        fail(f"verification-evidence.json is missing required fields: {missing}")
    if evidence.get("schema_version") != 1:
        fail("verification-evidence.json schema_version must be 1")

    finding_slug = str(evidence.get("finding_slug") or "").strip()
    if not finding_slug:
        fail("verification-evidence.json finding_slug must not be empty")
    if finding is not None:
        expected_slug = str(finding.get("slug") or "").strip()
        if expected_slug and finding_slug != expected_slug:
            fail(
                "verification-evidence.json finding_slug does not match findings.json: "
                f"expected {expected_slug}, got {finding_slug}"
            )

    status = str(evidence.get("verification_status") or "").strip()
    if status not in ALLOWED_VERIFICATION_STATUSES:
        fail(f"verification-evidence.json verification_status is invalid: {status or '<missing>'}")
    if status != "confirmed_in_docker":
        fail(
            "confirmed bundles under confirmed/ require verification_status=confirmed_in_docker; "
            f"got {status}"
        )

    if evidence.get("docker_required") is not True:
        fail("verification-evidence.json docker_required must be true for confirmed bundles")
    if evidence.get("severity_escalation_attempted") is not True:
        fail("verification-evidence.json severity_escalation_attempted must be true for confirmed bundles")

    for field in (
        "docker_image",
        "docker_command",
        "expected_observation",
        "observed_observation",
        "oracle_token",
        "severity_escalation_result",
    ):
        value = str(evidence.get(field) or "").strip()
        if not value:
            fail(f"verification-evidence.json {field} must not be empty")
        if value in PLACEHOLDER_VERIFICATION_VALUES:
            fail(f"verification-evidence.json {field} must not use placeholder text: {value}")

    validate_bundle_relative_file(evidence.get("poc_path"), bundle_dir, "poc_path")
    evidence_files = evidence.get("evidence_files")
    if not isinstance(evidence_files, list) or not evidence_files:
        fail("verification-evidence.json evidence_files must be a non-empty list")
    seen: set[str] = set()
    for index, item in enumerate(evidence_files, start=1):
        rel = validate_bundle_relative_file(item, bundle_dir, f"evidence_files[{index}]")
        if rel in seen:
            fail(f"verification-evidence.json evidence_files contains duplicate path: {rel}")
        seen.add(rel)
    return evidence


def validate_attachment_note(note_path: Path, language: str) -> None:
    text = note_path.read_text(encoding="utf-8")
    if language == "zh-CN":
        if "附件目录说明" not in text:
            fail("attachment note is missing the expected Chinese title")
        if "说明：" in text:
            fail("attachment note must not contain a standalone 说明 field")
        if "- 原始路径：" not in text or "- 用途：" not in text:
            fail("Chinese attachment note must contain 原始路径 and 用途 fields")
    else:
        if "Attachment Index" not in text:
            fail("attachment note is missing the expected English title")
        if "Note:" in text or "- Note:" in text:
            fail("English attachment note must not contain a standalone Note field")
        if "- Original Path:" not in text or "- Purpose:" not in text:
            fail("English attachment note must contain Original Path and Purpose fields")
    validate_no_absolute_paths(text)
    validate_relative_attachment_refs(text, note_path.parent)


def warn_attachment_hygiene(bundle_dir: Path, texts: list[str], evidence: dict[str, object]) -> None:
    refs: set[str] = set()
    for text in texts:
        refs.update(re.findall(r"attachments/[^\s`)>]+", text))
    evidence_files = evidence.get("evidence_files")
    if isinstance(evidence_files, list):
        refs.update(str(item).strip() for item in evidence_files if str(item).strip().startswith("attachments/"))
    poc_path = str(evidence.get("poc_path") or "").strip()
    if poc_path.startswith("attachments/"):
        refs.add(poc_path)

    nested_refs = sorted(ref for ref in refs if re.search(r"attachments/(?:[^/]+/)*security-research-\d{8}-\d{6}", ref))
    if nested_refs:
        warn(
            "nested workspace attachment paths detected; prefer stable bundle-local attachments/<short-name> paths: "
            + ", ".join(nested_refs[:5])
        )

    attachments_dir = bundle_dir / "attachments"
    basename_map: dict[str, list[str]] = {}
    if attachments_dir.exists():
        for path in attachments_dir.rglob("*"):
            if path.is_file():
                rel = path.relative_to(bundle_dir).as_posix()
                basename_map.setdefault(path.name, []).append(rel)
    duplicate_names = {
        name: sorted(paths)
        for name, paths in basename_map.items()
        if len(paths) > 1
    }
    if duplicate_names:
        preview = "; ".join(f"{name}: {', '.join(paths[:3])}" for name, paths in sorted(duplicate_names.items())[:5])
        warn(f"duplicate attachment basenames detected; consider deduping flat/nested evidence copies: {preview}")


def validate_reproduction_supplement(supplement_path: Path, language: str) -> str:
    text = supplement_path.read_text(encoding="utf-8")
    if supplement_path.name in GENERIC_SUPPLEMENT_FILENAMES:
        fail(f"generic reproduction supplement filename is not allowed: {supplement_path.name}")
    if language == "zh-CN":
        required_markers = ["补充复现说明", "关键成功证据", "结论"]
    else:
        required_markers = ["Reproduction Supplement", "Key Success Evidence", "Conclusion"]
    missing = [marker for marker in required_markers if marker not in text]
    if missing:
        fail(f"reproduction supplement is missing required sections: {missing}")
    validate_no_absolute_paths(text)
    validate_relative_attachment_refs(text, supplement_path.parent)
    validate_language_consistency(
        text,
        f"reproduction supplement {supplement_path.name}",
        language,
        expected_markers=ZH_SUPPLEMENT_MARKERS if language == "zh-CN" else EN_SUPPLEMENT_MARKERS,
        opposite_markers=EN_SUPPLEMENT_MARKERS if language == "zh-CN" else ZH_SUPPLEMENT_MARKERS,
    )
    return text


def validate_relative_attachment_refs(text: str, bundle_dir: Path | None = None) -> None:
    for match in re.findall(r"attachments/[^\s`)>]+", text):
        if match.startswith("/"):
            fail(f"attachment reference must be bundle-relative, got: {match}")
        if bundle_dir is not None:
            resolved = (bundle_dir / match).resolve()
            if not resolved.exists():
                fail(f"attachment reference does not exist inside bundle: {match}")


def validate_recording_step_labels(script_path: Path, text: str) -> None:
    for pattern in RECORDING_ZERO_STEP_PATTERNS:
        match = pattern.search(text)
        if match:
            fail(
                f"bundle root recording script must not use 0/N step labels: "
                f"{script_path.name} contains {match.group(0)!r}. Use [代码]/[Code], [证据]/[Evidence], "
                "or another non-numeric label for non-reproduction sections."
            )


def shell_syntax_command(script_path: Path) -> list[str]:
    first_line = ""
    try:
        first_line = script_path.read_text(encoding="utf-8", errors="ignore").splitlines()[0]
    except IndexError:
        first_line = ""
    lowered = first_line.lower()
    if "bash" in lowered and shutil.which("bash"):
        return ["bash", "-n", str(script_path)]
    if re.search(r"\b(?:sh|dash)\b", lowered) and shutil.which("sh"):
        return ["sh", "-n", str(script_path)]
    if "zsh" in lowered and shutil.which("zsh"):
        return ["zsh", "-n", str(script_path)]
    if shutil.which("bash"):
        return ["bash", "-n", str(script_path)]
    if shutil.which("sh"):
        return ["sh", "-n", str(script_path)]
    fail("cannot perform shell syntax check because neither bash nor sh is available")


def validate_shell_syntax(script_path: Path) -> None:
    command = shell_syntax_command(script_path)
    proc = subprocess.run(command, capture_output=True, text=True)
    if proc.returncode != 0:
        output = ((proc.stdout or "") + (proc.stderr or "")).strip()
        fail(f"shell syntax check failed for {script_path.name}: {output}")


def line_prints_final_confirmation(line: str) -> bool:
    if not any(pattern.search(line) for pattern in SHELL_CONFIRMATION_PHRASE_PATTERNS):
        return False
    return any(pattern.search(line) for pattern in SHELL_CONFIRMATION_PRINT_PATTERNS)


def has_preceding_success_oracle(lines: list[str], index: int) -> bool:
    window = "\n".join(lines[max(0, index - 80): index + 1])
    return any(pattern.search(window) for pattern in SHELL_SUCCESS_ORACLE_PATTERNS)


def logical_shell_lines(text: str) -> list[str]:
    logical: list[str] = []
    current = ""
    for raw in text.splitlines():
        stripped = raw.strip()
        if not stripped or stripped.startswith("#"):
            continue
        if current:
            current = f"{current} {stripped}"
        else:
            current = stripped
        if current.endswith("\\"):
            current = current[:-1].rstrip()
            continue
        logical.append(current)
        current = ""
    if current:
        logical.append(current)
    return logical


def line_softens_success_oracle_failure(line: str) -> bool:
    if not SHELL_FAIL_OPEN_ORACLE_COMMAND_PATTERN.search(line):
        return False
    if not SHELL_FAIL_OPEN_SOFTENER_PATTERN.search(line):
        return False
    return not SHELL_FAIL_CLOSED_SOFTENER_PATTERN.search(line)


def validate_fail_open_success_oracles(script_path: Path, text: str) -> None:
    lines = logical_shell_lines(text)
    confirmation_indexes = [
        index for index, line in enumerate(lines)
        if line_prints_final_confirmation(line)
    ]
    if not confirmation_indexes:
        return
    first_confirmation = min(confirmation_indexes)
    for index, line in enumerate(lines[: first_confirmation + 1]):
        if not line_softens_success_oracle_failure(line):
            continue
        preview = line if len(line) <= 180 else line[:177] + "..."
        fail(
            f"{script_path.name} softens a success-oracle failure before final confirmation: {preview!r}. "
            "Do not use `|| echo`, `|| printf`, or `|| true` for grep/jq/curl/docker-log success checks "
            "before printing VULNERABILITY CONFIRMED, ATTACK SUCCESS, 漏洞已确认, or 攻击成功; fail closed with exit 1."
        )


def validate_shell_success_oracles(script_path: Path, text: str) -> None:
    validate_fail_open_success_oracles(script_path, text)
    lines = text.splitlines()
    for index, line in enumerate(lines):
        if not line_prints_final_confirmation(line):
            continue
        if has_preceding_success_oracle(lines, index):
            continue
        fail(
            f"{script_path.relative_to(script_path.parent)} prints a final vulnerability confirmation without a "
            "nearby or preceding concrete success oracle. Add a jq -e/grep -q/test/if/case/status/JSON/content "
            "assertion before printing VULNERABILITY CONFIRMED, ATTACK SUCCESS, 漏洞已确认, or 攻击成功."
        )


def resolve_script_parent_traversal_token(script_path: Path, bundle_dir: Path, token: str) -> Path:
    cleaned = token.strip().strip("'\"`);,")
    base = script_path.parent
    suffix = cleaned
    var_bases = {
        "$SCRIPT_DIR/": script_path.parent,
        "${SCRIPT_DIR}/": script_path.parent,
        "$BUNDLE_ROOT/": bundle_dir,
        "${BUNDLE_ROOT}/": bundle_dir,
        "$ATTACH_DIR/": bundle_dir / "attachments",
        "${ATTACH_DIR}/": bundle_dir / "attachments",
        "$PWD/": script_path.parent,
        "${PWD}/": script_path.parent,
    }
    for prefix, candidate_base in var_bases.items():
        if cleaned.startswith(prefix):
            base = candidate_base
            suffix = cleaned[len(prefix):]
            break
    return (base / suffix).resolve()


def line_is_display_only_shell(line: str) -> bool:
    stripped = line.strip()
    return bool(
        re.match(r"^(?:echo|printf|highlight_success|highlight_danger|highlight_note|focus_line|print_banner)\b", stripped)
    )


def line_can_depend_on_filesystem_path(line: str) -> bool:
    stripped = line.strip()
    if re.search(r"\b(?:cd|pushd|source|mkdir|cp|mv|rsync|tar|chmod|chown|ln|test|find)\b", stripped):
        return True
    if re.search(r"\b(?:bash|sh|zsh|python3?|node|ruby|php)\s+", stripped):
        return True
    if re.search(r"\bdocker\b.*(?:-v|--volume|--mount)", stripped):
        return True
    if re.search(r"\b(?:BUNDLE_ROOT|REPO_ROOT|PROJECT_ROOT|SOURCE_ROOT|TARGET_ROOT|WORKSPACE_ROOT)\s*=", stripped):
        return True
    return False


def validate_script_bundle_portability(script_path: Path, bundle_dir: Path, text: str) -> None:
    bundle_root = bundle_dir.resolve()
    for line in logical_shell_lines(text):
        if line_is_display_only_shell(line):
            continue
        if not line_can_depend_on_filesystem_path(line):
            continue
        for match in SCRIPT_PARENT_TRAVERSAL_PATTERN.finditer(line):
            token = match.group("token")
            if not token:
                continue
            resolved = resolve_script_parent_traversal_token(script_path, bundle_dir, token)
            try:
                resolved.relative_to(bundle_root)
            except ValueError:
                rel = script_path.relative_to(bundle_dir).as_posix()
                fail(
                    f"{rel} depends on a parent path outside the confirmed bundle: {token}. "
                    "Final bundle scripts must use SCRIPT_DIR, BUNDLE_ROOT, or ATTACH_DIR paths that stay within the bundle."
                )

    for line in logical_shell_lines(text):
        if line_is_display_only_shell(line):
            continue
        if not PACKAGE_MANAGER_INSTALL_PATTERN.search(line):
            continue
        if INSTALL_LIFECYCLE_SAFE_PATTERN.search(line):
            continue
        rel = script_path.relative_to(bundle_dir).as_posix()
        warn(
            f"{rel} runs a package manager install command without --ignore-scripts, offline/local fixture flags, "
            "or an obvious vendored dependency path; document why network/lifecycle-script noise is unavoidable"
        )


def validate_bundle_root_scripts(bundle_dir: Path, note_text: str) -> list[str]:
    script_paths = sorted(
        [
            path for path in bundle_dir.iterdir()
            if path.is_file() and path.suffix == ".sh" and path.name.startswith("run-")
        ]
    )
    for script_path in script_paths:
        text = script_path.read_text(encoding="utf-8")
        validate_shell_syntax(script_path)
        validate_no_absolute_paths(text)
        validate_script_bundle_portability(script_path, bundle_dir, text)
        validate_recording_step_labels(script_path, text)
        validate_shell_success_oracles(script_path, text)
        if script_path.name not in note_text:
            fail(f"bundle root script is missing from attachment note: {script_path.name}")
        if not os.access(script_path, os.X_OK):
            fail(f"bundle root script is not executable: {script_path.name}")
        if "docker" not in text.lower():
            fail(f"bundle root script must explicitly use Docker or Docker Compose: {script_path.name}")
        for pattern in HOST_FALLBACK_PATTERNS:
            if pattern.search(text):
                fail(f"bundle root script must not include host fallback execution modes: {script_path.name}")
        if not any(pattern.search(text) for pattern in SCRIPT_PAUSE_PATTERNS):
            fail(f"bundle root script must include reviewer-friendly pauses around key checkpoints: {script_path.name}")
        pause_hits = sum(len(pattern.findall(text)) for pattern in SCRIPT_RUNTIME_PAUSE_PATTERNS)
        if pause_hits < 3:
            fail(
                f"bundle root script must include multiple short pauses at key reviewer checkpoints, not just a single pause: {script_path.name}"
            )
        if not any(pattern.search(text) for pattern in SCRIPT_COLOR_PATTERNS):
            fail(f"bundle root script must include ANSI color or terminal highlighting support: {script_path.name}")
        if not any(pattern.search(text) for pattern in SCRIPT_STEP_PATTERNS):
            fail(f"bundle root script must include visible step markers or banner helpers: {script_path.name}")
    if not script_paths:
        fail("confirmed bundle must include at least one bundle-root run-*.sh reproduction helper")
    return [path.name for path in script_paths]


def is_shell_script(path: Path) -> bool:
    if path.suffix in {".sh", ".bash"}:
        return True
    try:
        first = path.read_text(encoding="utf-8", errors="ignore").splitlines()[0]
    except IndexError:
        return False
    return bool(re.search(r"^#!.*\b(?:sh|bash|dash|zsh)\b", first))


def validate_attachment_shell_scripts(attachments_dir: Path, bundle_dir: Path) -> None:
    for path in sorted(item for item in attachments_dir.rglob("*") if item.is_file()):
        if not is_shell_script(path):
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        validate_shell_syntax(path)
        validate_no_absolute_paths(text)
        validate_script_bundle_portability(path, bundle_dir, text)
        validate_shell_success_oracles(path, text)


def yaml_indent(line: str) -> int:
    return len(line) - len(line.lstrip(" "))


def strip_inline_yaml_comment(value: str) -> str:
    text = value.strip()
    if not text:
        return ""
    quote: str | None = None
    escaped = False
    for index, char in enumerate(text):
        if escaped:
            escaped = False
            continue
        if char == "\\":
            escaped = True
            continue
        if quote:
            if char == quote:
                quote = None
            continue
        if char in {"'", '"'}:
            quote = char
            continue
        if char == "#":
            if index == 0 or text[index - 1].isspace():
                return text[:index].strip()
    return text


def unquote_yaml_scalar(value: str) -> str:
    text = strip_inline_yaml_comment(value).strip()
    if not text:
        return ""
    if (text.startswith("'") and text.endswith("'")) or (text.startswith('"') and text.endswith('"')):
        return text[1:-1]
    return text


def split_inline_yaml_list(value: str) -> list[str]:
    text = strip_inline_yaml_comment(value).strip()
    if not (text.startswith("[") and text.endswith("]")):
        return []
    inner = text[1:-1].strip()
    if not inner:
        return []
    items: list[str] = []
    current = ""
    quote: str | None = None
    escaped = False
    for char in inner:
        if escaped:
            current += char
            escaped = False
            continue
        if char == "\\":
            current += char
            escaped = True
            continue
        if quote:
            current += char
            if char == quote:
                quote = None
            continue
        if char in {"'", '"'}:
            current += char
            quote = char
            continue
        if char == ",":
            items.append(unquote_yaml_scalar(current))
            current = ""
            continue
        current += char
    if current.strip():
        items.append(unquote_yaml_scalar(current))
    return [item for item in items if item]


def service_key_indent(lines: list[str]) -> tuple[int | None, int | None]:
    services_indent = None
    service_indent = None
    for raw in lines:
        stripped = strip_inline_yaml_comment(raw).strip()
        if not stripped or stripped.startswith("#"):
            continue
        indent = yaml_indent(raw)
        if stripped == "services:":
            services_indent = indent
            continue
        if services_indent is not None and indent > services_indent and re.match(r"^[A-Za-z0-9_.-]+:\s*$", stripped):
            service_indent = indent
            break
    return services_indent, service_indent


def line_is_under_service(raw: str, services_indent: int | None, service_indent: int | None) -> bool:
    if services_indent is None or service_indent is None:
        return True
    indent = yaml_indent(raw)
    return indent > service_indent


def collect_env_file_entries(lines: list[str]) -> list[tuple[int, str]]:
    services_indent, service_indent = service_key_indent(lines)
    entries: list[tuple[int, str]] = []
    index = 0
    while index < len(lines):
        raw = lines[index]
        stripped = strip_inline_yaml_comment(raw).strip()
        indent = yaml_indent(raw)
        match = re.match(r"^env_file\s*:\s*(.*)$", stripped)
        if not match or not line_is_under_service(raw, services_indent, service_indent):
            index += 1
            continue
        value = match.group(1).strip()
        if value:
            inline_items = split_inline_yaml_list(value)
            if inline_items:
                entries.extend((index + 1, item) for item in inline_items)
            elif value.startswith("{") and "path:" in value:
                path_match = re.search(r"path\s*:\s*([^,}]+)", value)
                if path_match:
                    entries.append((index + 1, unquote_yaml_scalar(path_match.group(1))))
            else:
                entries.append((index + 1, unquote_yaml_scalar(value)))
            index += 1
            continue
        block_indent = indent
        index += 1
        while index < len(lines):
            nested_raw = lines[index]
            nested = strip_inline_yaml_comment(nested_raw).strip()
            nested_indent = yaml_indent(nested_raw)
            if nested and nested_indent <= block_indent:
                break
            list_match = re.match(r"^-\s*(.*)$", nested)
            if list_match:
                item_value = list_match.group(1).strip()
                if item_value.startswith("{") and "path:" in item_value:
                    path_match = re.search(r"path\s*:\s*([^,}]+)", item_value)
                    if path_match:
                        entries.append((index + 1, unquote_yaml_scalar(path_match.group(1))))
                elif item_value.startswith("path:"):
                    entries.append((index + 1, unquote_yaml_scalar(item_value.split(":", 1)[1])))
                elif item_value:
                    entries.append((index + 1, unquote_yaml_scalar(item_value)))
            elif nested.startswith("path:"):
                entries.append((index + 1, unquote_yaml_scalar(nested.split(":", 1)[1])))
            index += 1
    return entries


def collect_volume_entries(lines: list[str]) -> list[tuple[int, str, str]]:
    services_indent, service_indent = service_key_indent(lines)
    entries: list[tuple[int, str, str]] = []
    index = 0
    while index < len(lines):
        raw = lines[index]
        stripped = strip_inline_yaml_comment(raw).strip()
        indent = yaml_indent(raw)
        if stripped != "volumes:" or not line_is_under_service(raw, services_indent, service_indent):
            index += 1
            continue
        block_indent = indent
        index += 1
        while index < len(lines):
            nested_raw = lines[index]
            nested = strip_inline_yaml_comment(nested_raw).strip()
            nested_indent = yaml_indent(nested_raw)
            if nested and nested_indent <= block_indent:
                break
            list_match = re.match(r"^-\s*(.*)$", nested)
            if not list_match:
                index += 1
                continue
            item_start = index
            item_value = list_match.group(1).strip()
            if item_value and not re.match(r"^[A-Za-z_][A-Za-z0-9_-]*\s*:", item_value):
                entries.append((item_start + 1, "short", unquote_yaml_scalar(item_value)))
                index += 1
                continue
            item_type = ""
            item_source = ""
            if item_value.startswith("{"):
                type_match = re.search(r"type\s*:\s*([^,}]+)", item_value)
                source_match = re.search(r"(?:source|src)\s*:\s*([^,}]+)", item_value)
                item_type = unquote_yaml_scalar(type_match.group(1)) if type_match else ""
                item_source = unquote_yaml_scalar(source_match.group(1)) if source_match else ""
            elif ":" in item_value:
                key, value = item_value.split(":", 1)
                if key.strip() == "type":
                    item_type = unquote_yaml_scalar(value)
                elif key.strip() in {"source", "src"}:
                    item_source = unquote_yaml_scalar(value)
            index += 1
            while index < len(lines):
                child_raw = lines[index]
                child = strip_inline_yaml_comment(child_raw).strip()
                child_indent = yaml_indent(child_raw)
                if child and child_indent <= nested_indent:
                    break
                if child.startswith("- "):
                    break
                if ":" in child:
                    key, value = child.split(":", 1)
                    key = key.strip()
                    if key == "type":
                        item_type = unquote_yaml_scalar(value)
                    elif key in {"source", "src"}:
                        item_source = unquote_yaml_scalar(value)
                index += 1
            if item_source:
                entries.append((item_start + 1, item_type or "long", item_source))
    return entries


def is_absolute_host_path(value: str) -> bool:
    return Path(value).is_absolute() or value.startswith("~") or bool(re.match(r"^[A-Za-z]:[\\/]", value))


def is_pathish_volume_source(value: str, kind: str) -> bool:
    if kind == "bind":
        return True
    return (
        value in {".", ".."}
        or value.startswith("./")
        or value.startswith("../")
        or value.startswith("/")
        or value.startswith("~")
        or "\\" in value
        or "/" in value
    )


def short_volume_source(spec: str) -> str:
    text = spec.strip()
    if not text or text.startswith(":"):
        return ""
    if re.match(r"^[A-Za-z]:[\\/]", text):
        second_colon = text.find(":", 2)
        return text if second_colon == -1 else text[:second_colon]
    return text.split(":", 1)[0] if ":" in text else ""


def validate_compose_local_path(compose_file: Path, bundle_dir: Path, value: str, label: str, *, must_be_file: bool = False) -> None:
    token = unquote_yaml_scalar(value)
    if not token:
        return
    if "$" in token:
        fail(f"{compose_file.name} {label} must use a static bundle-local path, got variable reference: {token}")
    if is_absolute_host_path(token):
        fail(f"{compose_file.name} {label} must not use an absolute host path: {token}")
    resolved = (compose_file.parent / token).resolve()
    try:
        resolved.relative_to(bundle_dir.resolve())
    except ValueError:
        fail(f"{compose_file.name} {label} escapes the confirmed bundle: {token}")
    if not resolved.exists():
        fail(f"{compose_file.name} {label} does not exist relative to the compose file: {token}")
    if must_be_file and not resolved.is_file():
        fail(f"{compose_file.name} {label} must point to a file: {token}")


def validate_compose_static_paths(compose_file: Path, bundle_dir: Path) -> None:
    lines = compose_file.read_text(encoding="utf-8", errors="ignore").splitlines()
    env_entries = collect_env_file_entries(lines)
    for line_number, env_path in env_entries:
        validate_compose_local_path(compose_file, bundle_dir, env_path, f"env_file on line {line_number}", must_be_file=True)

    for line_number, kind, value in collect_volume_entries(lines):
        source = short_volume_source(value) if kind == "short" else value
        if not source:
            continue
        if kind == "short" and not is_pathish_volume_source(source, kind):
            continue
        if kind not in {"short", "bind"} and not is_pathish_volume_source(source, kind):
            continue
        validate_compose_local_path(compose_file, bundle_dir, source, f"volume source on line {line_number}")


def docker_compose_available() -> bool:
    docker = shutil.which("docker")
    if not docker:
        return False
    proc = subprocess.run([docker, "compose", "version"], capture_output=True, text=True, timeout=15)
    return proc.returncode == 0


def validate_docker_compose_config(compose_file: Path) -> None:
    if not docker_compose_available():
        warn(f"Docker Compose config check skipped for {compose_file.name} because `docker compose` is unavailable.")
        return
    proc = subprocess.run(
        ["docker", "compose", "-f", str(compose_file), "config"],
        cwd=compose_file.parent,
        capture_output=True,
        text=True,
        timeout=30,
    )
    if proc.returncode != 0:
        output = ((proc.stdout or "") + (proc.stderr or "")).strip()
        fail(f"Docker Compose static config check failed for {compose_file.name}: {output}")


def validate_attachment_compose_files(attachments_dir: Path, bundle_dir: Path) -> None:
    for compose_file in sorted(path for path in attachments_dir.rglob("*") if path.is_file() and COMPOSE_FILE_PATTERN.match(path.name)):
        validate_no_absolute_paths(compose_file.read_text(encoding="utf-8", errors="ignore"))
        validate_compose_static_paths(compose_file, bundle_dir)
        validate_docker_compose_config(compose_file)


def split_markdown_paragraphs(text: str) -> list[str]:
    paragraphs: list[str] = []
    current: list[str] = []
    in_fence = False
    for raw in text.splitlines():
        stripped = raw.strip()
        if stripped.startswith("```") or stripped.startswith("~~~"):
            if current:
                paragraphs.append(" ".join(current).strip())
                current = []
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        if not stripped:
            if current:
                paragraphs.append(" ".join(current).strip())
                current = []
            continue
        current.append(stripped)
    if current:
        paragraphs.append(" ".join(current).strip())
    return paragraphs


def looks_like_technical_fragment(text: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return True
    if any(pattern.search(stripped) for pattern in TECHNICAL_TEXT_PATTERNS):
        return True
    punctuation = sum(1 for char in stripped if char in "/:_=|{}[]<>`$\\")
    if punctuation >= 4:
        return True
    return False


def looks_like_long_english_prose(text: str) -> bool:
    stripped = re.sub(r"^[#>*\-\s\d.、)）]+", "", text.strip())
    if not stripped:
        return False
    if any("\u4e00" <= ch <= "\u9fff" for ch in stripped):
        return False
    if looks_like_technical_fragment(stripped):
        return False
    words = re.findall(r"\b[A-Za-z]{2,}\b", stripped)
    if len(words) < 16:
        return False
    letters = sum(1 for ch in stripped if ch.isalpha())
    ascii_letters = sum(1 for ch in stripped if ("a" <= ch.lower() <= "z"))
    if letters and ascii_letters / letters < 0.85:
        return False
    if not re.search(r"[.!?]\s+[A-Z]", stripped) and len(words) < 24:
        return False
    natural_markers = {
        "the",
        "this",
        "that",
        "with",
        "without",
        "attacker",
        "server",
        "vulnerability",
        "endpoint",
        "request",
        "response",
        "because",
        "allows",
        "could",
        "should",
        "impact",
        "security",
    }
    hits = sum(1 for word in words if word.lower() in natural_markers)
    return hits >= 4


def validate_no_untranslated_english_text(lines: list[str], markdown_texts: list[tuple[str, str]], language: str) -> None:
    if language != "zh-CN":
        return
    for index, line in enumerate(lines, start=1):
        if looks_like_long_english_prose(line):
            fail(
                "zh-CN report contains a long English natural-language paragraph in DOCX "
                f"paragraph {index}; translate reviewer-facing prose and keep only technical tokens in English"
            )
    for label, text in markdown_texts:
        for para_index, paragraph in enumerate(split_markdown_paragraphs(text), start=1):
            if looks_like_long_english_prose(paragraph):
                fail(
                    f"zh-CN artifact {label} contains a long English natural-language paragraph "
                    f"near paragraph {para_index}; translate reviewer-facing prose and keep commands, URLs, "
                    "headers, JSON keys, CVSS vectors, and oracle tokens as technical text only"
                )


def explicit_structured_values(*containers: object, fields: tuple[str, ...]) -> dict[str, str]:
    values: dict[str, str] = {}
    for container in containers:
        if not isinstance(container, dict):
            continue
        for field in fields:
            value = str(container.get(field) or "").strip()
            if value:
                values[field] = value
        nested = container.get("verification_evidence")
        if isinstance(nested, dict):
            for field in fields:
                value = str(nested.get(field) or "").strip()
                if value:
                    values[field] = value
    return values


def validate_structured_material_consistency(
    bundle_dir: Path,
    finding: dict[str, object] | None,
    evidence: dict[str, object],
    docx_text_value: str,
    supplement_text: str,
    root_scripts: list[str],
) -> None:
    finding_obj = finding or {}
    target_values = explicit_structured_values(finding_obj, evidence, fields=STRUCTURED_TARGET_FIELDS)
    material_texts = [docx_text_value, supplement_text]
    for script_name in root_scripts:
        path = bundle_dir / script_name
        if path.exists():
            material_texts.append(path.read_text(encoding="utf-8", errors="ignore"))
    attachments_dir = bundle_dir / "attachments"
    if attachments_dir.exists():
        for path in sorted(attachments_dir.rglob("*")):
            if path.is_file() and is_shell_script(path):
                material_texts.append(path.read_text(encoding="utf-8", errors="ignore"))
    combined = "\n".join(material_texts)
    for field, value in target_values.items():
        if value not in combined:
            fail(
                f"structured {field} from findings/verification-evidence does not appear in reviewer-facing "
                "DOCX, supplement, PoC, or recording script material"
            )

    docker_command = str(evidence.get("docker_command") or "").strip()
    if docker_command:
        docker_material = "\n".join([docx_text_value, supplement_text] + [
            (bundle_dir / script_name).read_text(encoding="utf-8", errors="ignore")
            for script_name in root_scripts
            if (bundle_dir / script_name).exists()
        ])
        if docker_command not in docker_material:
            warn(
                "verification-evidence.json docker_command is not repeated exactly in the DOCX, "
                "supplement, or recording script; keep future generated artifacts aligned when possible"
            )


def validate_success_evidence(text: str, language: str, *, label: str = "report") -> None:
    if language == "zh-CN":
        required_anchors = ["结果证据："] if label == "report" else ["结果证据：", "关键成功证据"]
        success_signals = ["攻击成功", "成功读取", "成功获取", "VULNERABILITY CONFIRMED", "ATTACK SUCCESS", "实际结果："]
    else:
        required_anchors = ["Evidence:"] if label == "report" else ["Evidence:", "Key Success Evidence"]
        success_signals = [
            "attack success",
            "successfully read",
            "successful read",
            "successfully retrieved",
            "successful retrieval",
            "vulnerability confirmed",
            "observed result:",
            "sensitive-file read",
            "proof",
        ]
    lowered = text.lower()
    if not any(anchor in text for anchor in required_anchors):
        fail(f"{label} must contain an explicit success-evidence block: {' / '.join(required_anchors)}")
    if not any(signal.lower() in lowered for signal in success_signals):
        fail(f"{label} is missing a clear success oracle proving reproduction or exploitation succeeded")


def extract_poc_labels(text: str) -> set[int]:
    labels: set[int] = set()
    for match in POC_LABEL_PATTERN.finditer(text or ""):
        try:
            labels.add(int(match.group(1)))
        except ValueError:
            continue
    return labels


def warn_poc_label_consistency(
    bundle_dir: Path,
    docx_text_value: str,
    note_text: str,
    supplement_text: str,
    evidence: dict[str, object],
    root_scripts: list[str],
) -> None:
    evidence_text = json.dumps(evidence, ensure_ascii=False, sort_keys=True)
    reviewer_labels = extract_poc_labels("\n".join([docx_text_value, note_text, supplement_text, evidence_text]))
    if not reviewer_labels:
        return
    script_texts = []
    for script_name in root_scripts:
        script_path = bundle_dir / script_name
        if script_path.exists():
            script_texts.append(script_path.read_text(encoding="utf-8", errors="ignore"))
    script_labels = extract_poc_labels("\n".join(script_texts))
    highest_reviewer = max(reviewer_labels)
    highest_script = max(script_labels) if script_labels else 0
    if highest_reviewer > highest_script:
        warn(
            f"root recording script appears to miss the highest PoC label: reviewer-facing materials mention "
            f"PoC-{highest_reviewer}, but root scripts only mention up to "
            f"{'PoC-' + str(highest_script) if highest_script else 'no PoC labels'}"
        )


def warn_recording_video_staleness(bundle_dir: Path, material_paths: list[Path]) -> None:
    videos = sorted(
        path for path in bundle_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in VIDEO_SUFFIXES
    )
    if not videos:
        return
    newest_material = max(
        (path.stat().st_mtime for path in material_paths if path.exists() and path.is_file()),
        default=0.0,
    )
    for video in videos:
        if newest_material > video.stat().st_mtime + 1:
            warn(
                "recording appears older than current reproduction script or report material: "
                + video.relative_to(bundle_dir).as_posix()
            )


def run_command(command: list[str]) -> tuple[int, str]:
    proc = subprocess.run(command, capture_output=True, text=True)
    output = (proc.stdout or "") + (proc.stderr or "")
    return proc.returncode, output.strip()


def optional_file_and_unzip_checks(docx_path: Path) -> None:
    if shutil.which("file"):
        code, output = run_command(["file", str(docx_path)])
        if code != 0:
            fail(f"`file` check failed: {output}")
        if "OOXML" not in output and "Microsoft Word 2007+" not in output and "Zip archive data" not in output:
            fail(f"`file` output does not look like a Word OOXML document: {output}")
    if shutil.which("unzip"):
        code, output = run_command(["unzip", "-t", str(docx_path)])
        if code != 0:
            fail(f"`unzip -t` failed: {output}")


def optional_libreoffice_check(docx_path: Path) -> None:
    soffice = shutil.which("soffice")
    if not soffice:
        print("WARN: LibreOffice check skipped because `soffice` is not installed.")
        return
    with tempfile.TemporaryDirectory(prefix="asr-soffice-") as tempdir:
        command = [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            tempdir,
            str(docx_path),
        ]
        code, output = run_command(command)
        if code != 0:
            fail(f"LibreOffice conversion failed: {output}")
        generated = list(Path(tempdir).glob("*.pdf"))
        if not generated:
            fail("LibreOffice conversion did not produce a PDF output")


def optional_markitdown_check(docx_path: Path) -> None:
    executable = shutil.which("markitdown")
    if not executable:
        print("WARN: MarkItDown check skipped because `markitdown` is not installed.")
        return
    code, output = run_command([executable, str(docx_path)])
    if code != 0:
        fail(f"MarkItDown extraction failed: {output}")
    if not output.strip():
        fail("MarkItDown extraction returned empty output")


def main() -> None:
    args = parse_args()
    bundle_dir = Path(args.bundle_dir).expanduser().resolve()
    if not bundle_dir.exists() or not bundle_dir.is_dir():
        fail(f"bundle directory does not exist: {bundle_dir}")
    validate_bundle_cleanliness(bundle_dir)
    workspace_dir = validate_workspace_metadata(bundle_dir)

    docx_path = find_single([p for p in bundle_dir.glob("*.docx") if p.is_file()], "report docx")
    markdown_paths = [p for p in bundle_dir.glob("*.md") if p.is_file()]
    expected_note_name = f"{docx_path.stem}{'_attachment_index.md' if args.language == 'en-US' else '_附件目录说明.md'}" if args.language != "auto" else None
    if expected_note_name:
        attachment_candidates = [p for p in markdown_paths if p.name == expected_note_name]
    else:
        attachment_candidates = [p for p in markdown_paths if p.name.endswith("_附件目录说明.md") or p.name.endswith("_attachment_index.md")]
    note_path = find_single(attachment_candidates, "attachment note markdown")

    validate_docx_container(docx_path)
    optional_file_and_unzip_checks(docx_path)

    lines = docx_text(docx_path)
    if not lines:
        fail("report docx has no readable text paragraphs")

    language = detect_language(docx_path, lines, note_path) if args.language == "auto" else args.language
    validate_output_filenames(docx_path, note_path, language)
    supplement_name = expected_supplement_filename(docx_path, language)
    supplement_path = find_single([p for p in markdown_paths if p.name == supplement_name], "reproduction supplement markdown")
    validate_required_headings(lines, language)
    validate_quality_gate(lines, language)
    supplement_text = validate_reproduction_supplement(supplement_path, language)
    exploitability_text = validate_real_world_exploitability(lines, supplement_text, language)
    validate_report_depth(lines, language)
    validate_bundle_identity(bundle_dir, lines, workspace_dir)
    findings_path = resolve_findings_path(bundle_dir)
    project_name = ""
    title_tokens: list[str] = []
    selected_finding: dict[str, object] | None = None
    if findings_path is not None and findings_path.exists():
        defaults, selected_finding = load_selected_finding(findings_path, bundle_dir.name, workspace_dir)
        project_name, title_tokens, selected_finding = load_bundle_identity(findings_path, bundle_dir.name, workspace_dir)
        validate_finding_vulnerability_name(selected_finding)
        validate_severity_consistency(bundle_dir, docx_path, selected_finding, language)
        validate_report_body_severity_consistency(bundle_dir, docx_path, lines, selected_finding, language)
        validate_title_auth_cvss_consistency(bundle_dir, docx_path, lines, selected_finding)
    verification_evidence = validate_verification_evidence(bundle_dir, selected_finding)
    if selected_finding is not None:
        validate_runtime_scope(defaults, selected_finding, verification_evidence, workspace_dir)

    combined_text = "\n".join(lines)
    if "CVSS 2.0" in combined_text:
        fail("CVSS 2.0 is not allowed; use CVSS 4.0 by default or CVSS 3.1 when required")
    validate_no_absolute_paths(combined_text)
    validate_no_placeholder_text(combined_text, language, "report docx")
    validate_relative_attachment_refs(combined_text, bundle_dir)
    reviewer_addendum_text = validate_reviewer_evidence_addendum(bundle_dir, language)
    validate_attachment_note(note_path, language)
    note_text = note_path.read_text(encoding="utf-8")
    validate_no_untranslated_english_text(
        lines,
        [(path.name, path.read_text(encoding="utf-8")) for path in markdown_paths],
        language,
    )
    validate_no_placeholder_text(note_text, language, "attachment note")
    validate_no_placeholder_text(supplement_text, language, "reproduction supplement")
    root_scripts = validate_bundle_root_scripts(bundle_dir, note_text)
    if project_name:
        validate_material_identity(supplement_text, f"reproduction supplement {supplement_path.name}", project_name, title_tokens)
        for script_name in root_scripts:
            script_text = (bundle_dir / script_name).read_text(encoding="utf-8")
            validate_material_identity(script_text, f"bundle root script {script_name}", project_name, title_tokens)
            validate_language_consistency(
                script_text,
                f"bundle root script {script_name}",
                language,
                expected_markers=ZH_REVIEW_MARKERS if language == "zh-CN" else EN_REVIEW_MARKERS,
                opposite_markers=EN_REVIEW_MARKERS if language == "zh-CN" else ZH_REVIEW_MARKERS,
            )
    validate_success_evidence(combined_text, language, label="report")
    validate_success_evidence(supplement_text, language, label="reproduction supplement")

    attachments_dir = bundle_dir / "attachments"
    if not attachments_dir.exists():
        fail("confirmed bundle must include an attachments/ directory, even when the reproduction helper is at bundle root")
    if not attachments_dir.is_dir():
        fail("attachments exists but is not a directory")
    if not any(path.is_file() for path in attachments_dir.rglob("*")):
        fail("attachments/ must contain at least one evidence, PoC, Docker, or supporting file")
    validate_attachment_shell_scripts(attachments_dir, bundle_dir)
    validate_attachment_compose_files(attachments_dir, bundle_dir)
    attachment_script_texts = [
        path.read_text(encoding="utf-8", errors="ignore")
        for path in sorted(attachments_dir.rglob("*"))
        if path.is_file() and is_shell_script(path)
    ]
    root_script_paths = [bundle_dir / script_name for script_name in root_scripts]
    root_script_texts = [
        path.read_text(encoding="utf-8", errors="ignore")
        for path in root_script_paths
        if path.exists()
    ]
    reviewer_index_text, _reviewer_index = validate_reviewer_evidence_index(
        bundle_dir,
        verification_evidence,
        combined_text,
        note_text,
        supplement_text,
        reviewer_addendum_text,
    )
    reviewer_story_text = "\n".join(
        [
            combined_text,
            note_text,
            supplement_text,
            reviewer_addendum_text,
            reviewer_index_text,
            json.dumps(verification_evidence, ensure_ascii=False, sort_keys=True),
            *root_script_texts,
            *attachment_script_texts,
        ]
    )
    validate_reviewer_story_gates(reviewer_story_text)
    validate_claim_oracle_consistency(reviewer_story_text)
    validate_poc_attacker_capability_boundary(
        "\n".join(
            [
                combined_text,
                note_text,
                supplement_text,
                reviewer_addendum_text,
                reviewer_index_text,
                json.dumps(verification_evidence, ensure_ascii=False, sort_keys=True),
                *root_script_texts,
                *attachment_script_texts,
            ]
        ),
        exploitability_text,
        language,
    )
    warn_poc_label_consistency(
        bundle_dir,
        combined_text,
        note_text,
        supplement_text,
        verification_evidence,
        root_scripts,
    )
    warn_recording_video_staleness(
        bundle_dir,
        [docx_path, note_path, supplement_path, bundle_dir / "verification-evidence.json", *root_script_paths],
    )
    validate_structured_material_consistency(
        bundle_dir,
        selected_finding,
        verification_evidence,
        combined_text,
        supplement_text,
        root_scripts,
    )
    warn_attachment_hygiene(bundle_dir, [combined_text, note_text, supplement_text], verification_evidence)

    if args.with_libreoffice:
        optional_libreoffice_check(docx_path)
    if args.with_markitdown:
        optional_markitdown_check(docx_path)

    if args.write_audit_event:
        write_bundle_validated_event(
            workspace_dir,
            bundle_dir,
            language,
            docx_path,
            note_path,
            supplement_path,
            verification_evidence,
        )

    print(f"VALIDATION PASSED: {bundle_dir}")
    print(f"language={language}")
    print(f"docx={docx_path.name}")
    print(f"attachment_note={note_path.name}")
    print(f"reproduction_supplement={supplement_path.name}")
    print(f"attachments_present={'yes' if attachments_dir.exists() else 'no'}")
    print(f"bundle_root_scripts={','.join(root_scripts) if root_scripts else 'none'}")


if __name__ == "__main__":
    main()
