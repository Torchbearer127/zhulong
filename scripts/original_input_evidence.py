#!/usr/bin/env python3
"""Bind supplied original-input artifacts; hashes do not attest capture truth."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path, PurePosixPath
import re
import stat

MANIFEST = "attachments/original-input-evidence.json"
FIELDS = {"text_path", "screenshot_path", "title", "description", "caption"}


def invalid(message: str) -> ValueError:
    return ValueError("ORIGINAL_INPUT_INVALID: " + message)


def artifact(root: Path, value: str) -> Path:
    if not isinstance(value, str) or not value.startswith("attachments/") or "\\" in value:
        raise invalid("use a bundle-local attachment path")
    parts = PurePosixPath(value).parts
    if any(part in {"..", "."} for part in parts):
        raise invalid("unsafe attachment path")
    path = root
    try:
        for part in parts:
            path = path / part
            if path.is_symlink():
                raise invalid("symlinked attachment")
        info = path.stat()
        if not stat.S_ISREG(info.st_mode) or info.st_nlink != 1:
            raise invalid("attachment must be a regular unlinked file")
    except OSError as exc:
        raise invalid("missing attachment: " + value) from exc
    return path


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check_spec(root: Path, spec: dict, language: str) -> None:
    from PIL import Image
    if not isinstance(spec, dict) or set(spec) != FIELDS:
        raise invalid("declaration fields must be text_path, screenshot_path, title, description, caption")
    for key in ("title", "description", "caption"):
        value = spec[key]
        if not isinstance(value, str) or not value.strip():
            raise invalid("empty reviewer text")
        if language == "zh-CN" and (not re.search(r"[\u4e00-\u9fff]", value) or re.search(r"Figure\s*:|Review evidence|Full packet", value, re.I)):
            raise invalid("Chinese report requires Chinese original-input title, description and caption")
    raw = artifact(root, spec["text_path"])
    try:
        text = raw.read_text(encoding="utf-8")
    except (UnicodeError, OSError) as exc:
        raise invalid("original input must be UTF-8 text") from exc
    if not text.strip() or any(ord(char) < 32 and char not in "\t\r\n" for char in text) or "\x7f" in text or len(raw.read_bytes()) > 32768:
        raise invalid("original input must be nonempty displayable UTF-8 text of at most 32 KiB")
    screenshot = artifact(root, spec["screenshot_path"])
    try:
        with Image.open(screenshot) as image:
            if image.format != "PNG" or image.width < 1 or image.height < 1:
                raise invalid("original screenshot must be a PNG")
            image.verify()
    except (OSError, SyntaxError, ValueError, Image.DecompressionBombError) as exc:
        raise invalid("invalid original screenshot") from exc


def prepare_original_input(root: Path, spec: dict, language: str) -> dict:
    check_spec(root, spec, language)
    manifest = {"schema_version": 1, "language": language, **spec,
                "text_sha256": digest(artifact(root, spec["text_path"])),
                "screenshot_sha256": digest(artifact(root, spec["screenshot_path"]))}
    (root / MANIFEST).write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def render_original_input(doc, root: Path, manifest: dict) -> None:
    from docx.shared import Inches
    doc.add_heading(manifest["title"], level=2)
    doc.add_paragraph(manifest["description"])
    doc.add_paragraph(manifest["text_path"])
    doc.add_picture(str(artifact(root, manifest["screenshot_path"])), width=Inches(6))
    doc.add_paragraph(manifest["caption"])


def finding_declaration(root: Path) -> tuple[bool, object]:
    path = root / "findings.json"
    if not path.exists() and not path.is_symlink():
        return False, None
    try:
        info = path.lstat()
        if not stat.S_ISREG(info.st_mode) or info.st_nlink != 1:
            raise invalid("findings.json must be a regular unlinked file")
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, ValueError) as exc:
        raise invalid("cannot read findings.json") from exc
    records = payload.get("findings", [payload]) if isinstance(payload, dict) else payload
    if not isinstance(records, list) or not all(isinstance(item, dict) for item in records):
        raise invalid("invalid findings.json records")
    declared = [item["original_input"] for item in records if "original_input" in item]
    if not declared:
        return False, None
    if len(records) != 1:
        raise invalid("original input requires one unambiguous finding")
    return True, declared[0]


def validate_original_input(root: Path, *, required: bool = False, language: str | None = None) -> dict | None:
    from docx import Document
    from docx.oxml.ns import qn
    declared, spec = finding_declaration(root)
    path = root / MANIFEST
    if not path.exists() and not path.is_symlink():
        if required or declared:
            raise invalid("original input declaration is required")
        return None
    try:
        manifest = json.loads(artifact(root, MANIFEST).read_text(encoding="utf-8"))
    except (ValueError, OSError) as exc:
        raise invalid("cannot read declaration") from exc
    if not isinstance(manifest, dict) or set(manifest) != FIELDS | {"schema_version", "language", "text_sha256", "screenshot_sha256"}:
        raise invalid("invalid declaration fields")
    if not declared or not isinstance(spec, dict) or set(spec) != FIELDS or spec != {key: manifest[key] for key in FIELDS}:
        raise invalid("attachment binding differs from finding.original_input")
    if type(manifest["schema_version"]) is not int or manifest["schema_version"] != 1 or not isinstance(manifest["language"], str) or manifest["language"] not in {"zh-CN", "en-US"}:
        raise invalid("invalid declaration version/language")
    if language is not None and language != manifest["language"]:
        raise invalid("declaration language differs from report language")
    check_spec(root, {key: manifest[key] for key in FIELDS}, manifest["language"])
    for role in ("text", "screenshot"):
        if digest(artifact(root, manifest[role + "_path"])) != manifest[role + "_sha256"]:
            raise invalid(role + " hash mismatch")
    reports = list(root.glob("*.docx"))
    if len(reports) != 1:
        raise invalid("exactly one DOCX is required")
    try:
        info = reports[0].lstat()
        if not stat.S_ISREG(info.st_mode) or info.st_nlink != 1:
            raise invalid("DOCX must be a regular unlinked file")
        doc = Document(reports[0])
    except Exception as exc:
        raise invalid("cannot parse DOCX") from exc
    paragraphs = doc.paragraphs
    matched = False
    for index, paragraph in enumerate(paragraphs):
        for blip in paragraph._p.xpath(".//a:blip"):
            relation = doc.part.rels.get(blip.get(qn("r:embed")))
            if relation is None or relation.is_external:
                continue
            if hashlib.sha256(relation.target_part.blob).hexdigest() != manifest["screenshot_sha256"]:
                continue
            if index >= 3 and index + 1 < len(paragraphs) and [p.text for p in paragraphs[index - 3:index]] == [manifest["title"], manifest["description"], manifest["text_path"]] and paragraphs[index + 1].text == manifest["caption"]:
                matched = True
    if not matched:
        raise invalid("DOCX must embed the exact supplied screenshot beside its title, description, original-text path and caption")
    return manifest
